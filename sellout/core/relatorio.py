"""Relatório de conferência da rodada, em xlsx."""

from __future__ import annotations

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

CABECALHO_FONTE = Font(bold=True, color="FFFFFF")
CABECALHO_FUNDO = PatternFill("solid", fgColor="1F3864")


def _aba(wb, titulo, colunas, linhas, larguras):
    ws = wb.create_sheet(titulo[:31])
    ws.append(colunas)
    for c in ws[1]:
        c.font = CABECALHO_FONTE
        c.fill = CABECALHO_FUNDO
    for linha in linhas:
        ws.append(linha)
    for i, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


def gerar(rel: dict, destino: str) -> str:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    _aba(wb, "Resumo", ["Indicador", "Valor"], [
        ["Linhas atualizadas", rel.get("alterados", 0)],
        ["Produtos novos inseridos", len(rel.get("novos_inseridos", []))],
        ["Linhas divididas por cor", len(rel.get("por_cor", []))],
        ["Lançamentos de produção aplicados", len(rel.get("producao", []))],
        ["Fórmulas de bloco ajustadas", len(rel.get("formulas_ajustadas", []))],
        ["Linhas mantidas sem alteração", len(rel.get("mantidos", []))],
        ["Linhas marcadas para revisar", len(rel.get("pendentes", []))],
    ], [42, 16])

    _aba(wb, "Nivel de Estoque",
         ["Planilha", "Aba", "Peças", "Valor (R$)", "Preço médio (R$)"],
         [[pl, aba, d["pecas"], round(d["valor"], 2),
           round(d["valor"] / d["pecas"], 2) if d["pecas"] else 0]
          for pl, abas in rel.get("nivel", {}).items() for aba, d in abas.items()],
         [14, 12, 12, 16, 18])

    _aba(wb, "Novos produtos inseridos",
         ["Planilha", "Aba", "Bloco", "Linha", "Código", "Descrição"],
         [[d["planilha"], d["aba"], d["bloco"], d["linha"], d["codigo"], d["descricao"]]
          for d in rel.get("novos_inseridos", [])],
         [12, 12, 22, 8, 12, 45])

    _aba(wb, "Producao aplicada",
         ["Planilha", "Aba", "Linha", "Código", "Descrição", "Qtde", "Observação"],
         [[d["planilha"], d["aba"], d["linha"], d["codigo"], d["descricao"], d["qtde"], d.get("obs", "")]
          for d in rel.get("producao", [])],
         [12, 12, 8, 12, 45, 10, 38])

    nao_aplicada = [[d["codigo"], d["qtde"], d.get("motivo", "")]
                    for d in rel.get("producao_nao_aplicada", [])]
    nao_aplicada += [[d["codigo"], d["qtde"], "produto não incluído nesta rodada"]
                     for d in rel.get("producao_sem_destino", [])]
    _aba(wb, "Producao nao aplicada", ["Código", "Qtde", "Situação"], nao_aplicada, [14, 10, 48])

    _aba(wb, "Divisao por cor",
         ["Planilha", "Aba", "Linha", "Código", "Descrição", "Texto em vermelho", "Cores atribuídas"],
         [[d["planilha"], d["aba"], d["linha"], d["codigo"], d["descricao"],
           d["vermelho"] or "(sem vermelho)", ", ".join(d["cores"])]
          for d in rel.get("por_cor", [])],
         [12, 12, 8, 12, 45, 28, 45])

    _aba(wb, "Formulas ajustadas",
         ["Planilha", "Aba", "Linha", "Tipo", "Antes", "Depois"],
         [[d["planilha"], d["aba"], d["linha"], d["tipo"], d["antes"], d["depois"]]
          for d in rel.get("formulas_ajustadas", [])],
         [12, 12, 8, 20, 46, 46])

    pendencias = [[d["planilha"], d["aba"], "linha %s / %s" % (d["linha"], d["codigo"]),
                   d.get("descricao", ""), "mantido sem alteração"]
                  for d in rel.get("mantidos", [])]
    pendencias += [[d["planilha"], d["aba"], "linha %s / %s" % (d["linha"], d["codigo"]),
                    d.get("descricao", ""), d.get("motivo", "divisão por cor não identificada")]
                   for d in rel.get("pendentes", [])]
    pendencias += [[d["planilha"], d["aba"], "linha %s / %s" % (d["linha"], d["codigo"]),
                    d.get("descricao", ""), "código não existe em Estoque, Vendas nem Produção"]
                   for d in rel.get("nao_encontrados", [])]
    pendencias += [[d["planilha"], d["aba"], "linha %s / %s" % (d["linha"], d["codigo"]),
                    d.get("descricao", ""), "sem cores nesta linha — Estoque atual zerado, "
                    "o saldo do código ficou na linha irmã"]
                   for d in rel.get("sem_cores", [])]
    pendencias += [[d["planilha"], d["aba"], d["codigo"], "",
                    "cores %s (%s pçs) sem linha de destino" % (", ".join(d["cores"]), d["estoque"])]
                   for d in rel.get("cores_sem_destino", [])]
    pendencias += [[d["planilha"], d["aba"], "%s / cor %s / %s" % (d["codigo"], d["cor"], d["tamanho"]), "",
                    "sem preço exato na aba Preco — usado %s (R$ %s) para %s pç"
                    % (d["origem"], d["preco_usado"], d["qtde"])]
                   for d in rel.get("sem_preco", [])]
    _aba(wb, "Pendencias", ["Planilha", "Aba", "Ref.", "Descrição", "Situação"],
         pendencias, [12, 12, 28, 40, 62])

    wb.save(destino)
    return destino
