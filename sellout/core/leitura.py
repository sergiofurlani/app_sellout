"""Leitura das planilhas: fontes de dados e estrutura das abas de trabalho."""

from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field

import openpyxl
from openpyxl.cell.rich_text import CellRichText
from openpyxl.utils import get_column_letter

from .cores import match, nrm

VERMELHOS = {"FFFF0000", "00FF0000"}
ABAS_TRABALHO = ("Masculino", "Feminino")


def num(valor):
    """Quantidade ou preço vindo de célula que pode estar como texto.

    Devolve o número, 0 para célula vazia, e None quando o conteúdo não é
    numérico — aí quem chamou registra a linha e segue com zero, em vez de
    a rodada inteira parar por causa de uma célula.
    """
    if valor is None or isinstance(valor, bool):
        return 0
    if isinstance(valor, (int, float)):
        return valor
    t = str(valor).strip().replace("\xa0", "").replace(" ", "")
    if not t or t in {"-", "--", "#N/D", "#N/A"}:
        return 0
    if "," in t and "." in t:          # 1.234,56 → 1234.56
        t = t.replace(".", "").replace(",", ".")
    elif "," in t:                     # 1,5 → 1.5
        t = t.replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None


def norm_codigo(valor):
    """Códigos vêm ora como texto, ora como número. Uniformiza."""
    if valor is None:
        return None
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def texto(valor):
    """Texto puro de uma célula que pode ser rich text."""
    if isinstance(valor, CellRichText):
        return "".join(str(t) for t in valor)
    return valor


def vermelho(valor) -> str:
    """Só os trechos escritos em vermelho — é onde a cor do produto aparece."""
    if not isinstance(valor, CellRichText):
        return ""
    partes = []
    for t in valor:
        fonte = getattr(t, "font", None)
        cor = getattr(fonte, "color", None) if fonte else None
        if cor is not None and getattr(cor, "rgb", None) in VERMELHOS:
            partes.append(str(t))
    return "".join(partes).strip()


def colecao_do_bloco(titulo) -> str | None:
    """"MASCULINO - SS 27" -> "SS27"."""
    achado = re.search(r"\b(SS|AW)\s*(\d{2})\b", nrm(titulo))
    return achado.group(1) + achado.group(2) if achado else None


# --------------------------------------------------------------------------- #
# Fontes de dados (sempre vindas da planilha geral)
# --------------------------------------------------------------------------- #

@dataclass
class Fontes:
    estoque: dict = field(default_factory=dict)        # cod -> cor -> qtd
    linhas_estoque: list = field(default_factory=list)  # (cod, cor, codcor, tam, qtd)
    # Grão cru, para gravar no banco: uma entrada por linha da aba, com a
    # filial e o tamanho preservados. O resto do Fontes agrega por produto e
    # cor porque é o que a planilha precisa; o banco quer o que o arquivo
    # disse, sem filtro — filtrar filial é decisão de consulta (D12).
    linhas: dict = field(default_factory=lambda: {
        "estoque": [], "vendas": [], "producao": [], "preco": []})
    vendas: dict = field(default_factory=dict)          # cod -> cor -> filial -> qtd
    producao: dict = field(default_factory=dict)        # cod -> cor -> qtd
    preco: dict = field(default_factory=dict)           # (cod, codcor, tam) -> valor
    preco_cor: dict = field(default_factory=dict)
    preco_produto: dict = field(default_factory=dict)
    produtos: list = field(default_factory=list)
    filiais_estoque: list = field(default_factory=list)
    filiais_vendas: list = field(default_factory=list)
    valores_ignorados: list = field(default_factory=list)
    avisos: list = field(default_factory=list)
    colunas_usadas: dict = field(default_factory=dict)
    cores_conhecidas: set = field(default_factory=set)
    # Entradas vindas do ERP (evento 106, ELENA ES -> lojas), no formato
    # cod -> cor normalizada -> quantidade. Vazio quando o arquivo não veio:
    # aí vale a regra antiga (D6), e o relatório diz qual foi usada.
    entradas_erp: dict = field(default_factory=dict)

    def eh_cor_conhecida(self, nome) -> bool:
        """O nome aparece como cor em alguma aba de origem?

        Serve para separar "cor que zerou o estoque nesta semana" de "texto que
        não é cor": a primeira só contribui zero, a segunda trava a divisão.
        """
        return bool(match(nome, self.cores_conhecidas))

    def cores_do_codigo(self, codigo) -> set:
        disp = set(self.estoque.get(codigo, {}))
        disp |= set(self.vendas.get(codigo, {}))
        disp |= set(self.producao.get(codigo, {}))
        disp.discard("")
        return disp

    def preco_de(self, codigo, codcor, tamanho):
        """Preço exato; se faltar, cai para o mesmo código."""
        exato = self.preco.get((codigo, codcor, tamanho))
        if exato is not None:
            return exato, "exato"
        cor = self.preco_cor.get((codigo, codcor))
        if cor is not None:
            return cor, "mesma cor"
        produto = self.preco_produto.get(codigo)
        if produto is not None:
            return produto, "mesmo código"
        return None, "sem preço"


ABAS_FONTE = ("Estoque", "Vendas", "Producao", "Preco", "Produtos")

# Colunas das abas de origem, procuradas pelo cabeçalho da linha 1. Cada entrada
# é (nome interno, lista de cabeçalhos aceitos, posição de reserva, obrigatória).
# O cabeçalho manda; a posição só entra se o texto não for encontrado.
COLUNAS_ESTOQUE = [
    ("filial", ["FILIAL"], 1, False),
    ("codigo", ["CODIGO", "COD"], 2, True),
    ("codigo_cor", ["CODIGO COR", "COD COR"], 4, True),
    ("cor", ["COR"], 5, True),
    ("tamanho", ["TAMANHO", "TAM"], 6, True),
    ("qtde", ["ESTOQUE ATUAL", "ESTOQUE", "QTDE", "QUANTIDADE", "SALDO"], 7, True),
]
COLUNAS_VENDAS = [
    ("filial", ["FILIAL"], 1, False),
    ("codigo", ["CODIGO", "COD"], 2, True),
    ("codigo_cor", ["CODIGO COR", "COD COR"], 4, False),
    ("cor", ["COR"], 5, True),
    ("tamanho", ["TAM", "TAMANHO"], 6, False),
    ("qtde", ["QTDE", "QUANTIDADE", "QTD"], 7, True),
]
COLUNAS_PRODUCAO = [
    ("codigo", ["CODIGO", "COD"], 1, True),
    ("cor", ["COR"], 3, True),
    ("tamanho", ["TAMANHO", "TAM"], 4, False),
    ("qtde", ["QUANTIDADE", "QTDE", "QTD"], 5, True),
]
COLUNAS_PRECO = [
    ("codigo", ["PRODUTO", "CODIGO", "COD"], 1, True),
    ("codigo_cor", ["CODIGO COR", "COD COR"], 3, True),
    ("tamanho", ["TAMANHO", "TAM"], 5, True),
    ("preco", ["PRECO", "PRECO VENDA", "VALOR"], 6, True),
]
COLUNAS_PRODUTOS = [
    ("codigo", ["CODIGO", "COD"], 1, True),
    ("descricao", ["DESCRICAO"], 3, True),
    ("colecao", ["COL", "COLECAO"], 8, True),
    ("divisao", ["DIVISAO", "DEPARTAMENTO"], 12, True),
]


def amostra_da_coluna(ws, coluna, quantas=4, ate=400):
    """Primeiros valores preenchidos, para mostrar na tela de conferência."""
    vistos = []
    for r in range(2, min(ws.max_row, ate) + 1):
        v = ws.cell(r, coluna).value
        if v is not None and str(v).strip():
            vistos.append(str(v)[:18])
            if len(vistos) >= quantas:
                break
    return vistos


def fracao_numerica(ws, coluna, ate=300):
    """Quanto da coluna é número — usado para achar a coluna de quantidade."""
    total = numericos = 0
    for r in range(2, min(ws.max_row, ate) + 1):
        v = ws.cell(r, coluna).value
        if v is None or not str(v).strip():
            continue
        total += 1
        if num(v) is not None:
            numericos += 1
    return numericos / total if total else 0.0


def colunas_com_dado(ws, ate=200) -> int:
    """Até onde as linhas de dados vão, independentemente do cabeçalho."""
    ultima = 0
    for r in range(2, min(ws.max_row, ate) + 1):
        for c in range(ws.max_column, ultima, -1):
            v = ws.cell(r, c).value
            if v is not None and str(v).strip():
                ultima = max(ultima, c)
                break
    return ultima


def cabecalho_deslocado(ws) -> bool:
    """Coluna inserida sem mexer na linha 1 deixa dados além do último rótulo."""
    rotulos = max((c for c in range(1, ws.max_column + 1)
                   if ws.cell(1, c).value not in (None, "")), default=0)
    return colunas_com_dado(ws) > rotulos


def ajusta_coluna_numerica(ws, achadas, campo, rotulo, avisos):
    """Confere se a coluna do número é mesmo numérica; se não, procura à direita.

    É o que salva a rodada quando inserem uma coluna no meio do relatório de
    origem sem mexer no cabeçalho: o rótulo "ESTOQUE ATUAL" acaba em cima do
    tamanho, e a quantidade fica na coluna seguinte, sem rótulo nenhum.
    """
    col = achadas.get(campo)
    if not col:
        return
    if fracao_numerica(ws, col) >= 0.9:
        return
    for candidata in range(col + 1, colunas_com_dado(ws) + 1):
        if fracao_numerica(ws, candidata) >= 0.95:
            avisos.append(
                "Na aba %s o rótulo \"%s\" está na coluna %s, que não tem números "
                "(ex.: %s). Usei a coluna %s, que é numérica (ex.: %s). "
                "Provavelmente inseriram uma coluna sem ajustar o cabeçalho."
                % (ws.title, rotulo, get_column_letter(col),
                   ", ".join(amostra_da_coluna(ws, col, 3)) or "vazia",
                   get_column_letter(candidata),
                   ", ".join(amostra_da_coluna(ws, candidata, 3)) or "vazia"))
            achadas[campo] = candidata
            return
    avisos.append(
        "Na aba %s não encontrei uma coluna numérica para %s; ela entra como zero."
        % (ws.title, rotulo))


def valores_da_coluna(ws, coluna, ate=400) -> set:
    vistos = set()
    for r in range(2, min(ws.max_row, ate) + 1):
        v = ws.cell(r, coluna).value
        if v is not None and str(v).strip():
            vistos.add(norm_codigo(v))
    return vistos


def ajusta_por_vocabulario(ws, achadas, campo, vocabulario, rotulo, avisos):
    """Confere a coluna contra os valores que a aba Preco usa para o mesmo campo.

    O cruzamento com o preço é por Código + Código Cor + Tamanho; se o
    cabeçalho escorregou, o tamanho vira outra coisa e o Nível de Estoque sai
    errado sem dar erro nenhum. Comparar com o vocabulário conhecido pega isso.
    """
    col = achadas.get(campo)
    if not col or not vocabulario:
        return

    def cobertura(c):
        vals = valores_da_coluna(ws, c)
        return len(vals & vocabulario) / len(vals) if vals else 0.0

    if cobertura(col) >= 0.5:
        return
    melhor, nota = None, 0.0
    for candidata in range(1, colunas_com_dado(ws) + 1):
        atual = cobertura(candidata)
        if atual > nota:
            melhor, nota = candidata, atual
    if melhor and melhor != col and nota >= 0.8:
        avisos.append(
            "Na aba %s o %s estava sendo lido da coluna %s (ex.: %s), que não bate "
            "com os valores da aba Preco. Usei a coluna %s (ex.: %s)."
            % (ws.title, rotulo, get_column_letter(col),
               ", ".join(amostra_da_coluna(ws, col, 3)) or "vazia",
               get_column_letter(melhor), ", ".join(amostra_da_coluna(ws, melhor, 3))))
        achadas[campo] = melhor


class ColunaAusente(Exception):
    """Cabeçalho obrigatório não encontrado numa aba de origem."""


def colunas_da_fonte(ws, especificacao) -> dict:
    """Casa os cabeçalhos da linha 1 com os nomes internos das colunas."""
    cabecalhos = {}
    for c in range(1, min(ws.max_column, 40) + 1):
        titulo = nrm(texto(ws.cell(1, c).value))
        if titulo and titulo not in cabecalhos:
            cabecalhos[titulo] = c

    achadas = {}
    for nome, aceitos, reserva, obrigatoria in especificacao:
        coluna = next((cabecalhos[a] for a in aceitos if a in cabecalhos), None)
        if coluna is None:
            # Cabeçalho renomeado: aceita quem começa com o texto esperado.
            coluna = next((pos for a in aceitos
                           for titulo, pos in cabecalhos.items() if titulo.startswith(a)), None)
        if coluna is None and reserva <= ws.max_column:
            coluna = reserva
        if coluna is None and obrigatoria:
            raise ColunaAusente(
                "A aba %s está sem a coluna %s." % (ws.title, aceitos[0]))
        achadas[nome] = coluna
    return achadas


def abas_faltando(caminho) -> list[str]:
    wb = openpyxl.load_workbook(caminho, read_only=True)
    try:
        presentes = set(wb.sheetnames)
    finally:
        wb.close()
    faltando = [a for a in ABAS_FONTE if a not in presentes]
    faltando += [a for a in ABAS_TRABALHO if a not in presentes]
    return faltando


def carrega_entradas_erp(caminho) -> dict:
    """Lê o CSV que o coletor gera: codigo;codigo_cor;cor;quant.

    É a ponte entre as duas metades do sistema. O MN só responde dentro da
    rede da Egrey, então o app na nuvem nunca vai buscar isto sozinho: o
    coletor puxa lá e o arquivo sobe junto com as planilhas.

    Sem o arquivo, `entradas_erp` fica vazio e tudo segue como antes.
    """
    import csv

    entradas = defaultdict(lambda: defaultdict(float))
    with open(caminho, newline="", encoding="utf-8-sig") as f:
        amostra = f.read(4096)
        f.seek(0)
        try:
            dialeto = csv.Sniffer().sniff(amostra, delimiters=";,\t")
        except csv.Error:
            dialeto = csv.excel
            dialeto.delimiter = ";"
        for linha in csv.DictReader(f, dialect=dialeto):
            chaves = {k.strip().lower(): v for k, v in linha.items() if k}
            cod = norm_codigo(chaves.get("codigo"))
            if not cod:
                continue
            q = num(chaves.get("quant") or chaves.get("quantidade"))
            if not q:
                continue
            # o ERP manda "0308 - VERMELHO"; as abas de origem trazem só o
            # nome, e é por ele que a divisão por cor casa
            cor = str(chaves.get("cor") or "")
            if " - " in cor:
                cor = cor.split(" - ", 1)[1]
            entradas[cod][nrm(cor)] += q
    return {c: dict(v) for c, v in entradas.items()}


def carrega_fontes(caminho, filiais_estoque=None, colunas_forcadas=None,
                   entradas_erp=None) -> Fontes:
    """Lê Estoque, Vendas, Producao, Preco e Produtos da planilha geral.

    `filiais_estoque` restringe quais filiais entram na soma de estoque.
    None significa todas as que existirem no arquivo.

    `colunas_forcadas` sobrepõe a coluna detectada, no formato
    {"Estoque:qtde": 8} — é o que a tela de conferência manda quando o
    usuário corrige a coluna do número.
    """
    colunas_forcadas = colunas_forcadas or {}
    wb = openpyxl.load_workbook(caminho, data_only=True)
    f = Fontes()
    if entradas_erp:
        f.entradas_erp = (entradas_erp if isinstance(entradas_erp, dict)
                          else carrega_entradas_erp(entradas_erp))

    def qtd_de(ws, aba, linha, coluna, rotulo):
        """Lê um número tolerando texto; registra a célula quando não dá."""
        bruto = ws.cell(linha, coluna).value
        valor = num(bruto)
        if valor is None:
            f.valores_ignorados.append(
                {"aba": aba, "linha": linha, "coluna": rotulo, "valor": str(bruto)[:40]})
            return 0
        return valor

    def prepara(ws, especificacao, campo_num, rotulo_num):
        achadas = colunas_da_fonte(ws, especificacao)
        if cabecalho_deslocado(ws):
            f.avisos.append(
                "A aba %s tem dados além do último rótulo da linha 1 — sinal de "
                "coluna inserida sem ajustar o cabeçalho." % ws.title)
        ajusta_coluna_numerica(ws, achadas, campo_num, rotulo_num, f.avisos)
        forcada = colunas_forcadas.get("%s:%s" % (ws.title, campo_num))
        if forcada:
            achadas[campo_num] = int(forcada)
        f.colunas_usadas[ws.title] = {
            "campo": campo_num,
            "rotulo": rotulo_num,
            "coluna": achadas[campo_num],
            "escolhida": bool(forcada),
            "opcoes": [
                {"coluna": i,
                 "letra": get_column_letter(i),
                 "cabecalho": (texto(ws.cell(1, i).value) or "(sem cabeçalho)"),
                 "amostra": ", ".join(amostra_da_coluna(ws, i, 3)) or "vazia"}
                for i in range(1, colunas_com_dado(ws) + 1)
            ],
        }
        return achadas

    pre = wb["Preco"]
    c = prepara(pre, COLUNAS_PRECO, "preco", "Preco")
    for r in range(2, pre.max_row + 1):
        cod = norm_codigo(pre.cell(r, c["codigo"]).value)
        codcor = norm_codigo(pre.cell(r, c["codigo_cor"]).value)
        tam = norm_codigo(pre.cell(r, c["tamanho"]).value)
        valor = qtd_de(pre, "Preco", r, c["preco"], "Preco")
        f.preco[(cod, codcor, tam)] = valor
        f.preco_cor.setdefault((cod, codcor), valor)
        f.preco_produto.setdefault(cod, valor)
        if cod:
            f.linhas["preco"].append(
                {"codigo": cod, "codigo_cor": codcor, "tamanho": tam, "preco": valor})

    vocab_tamanho = {t for (_c, _cc, t) in f.preco}
    vocab_codigo_cor = {cc for (_c, cc, _t) in f.preco}

    est = wb["Estoque"]
    c = prepara(est, COLUNAS_ESTOQUE, "qtde", "ESTOQUE ATUAL")
    ajusta_por_vocabulario(est, c, "tamanho", vocab_tamanho, "tamanho", f.avisos)
    ajusta_por_vocabulario(est, c, "codigo_cor", vocab_codigo_cor, "código da cor", f.avisos)
    f.estoque = defaultdict(lambda: defaultdict(int))
    vistas = []
    for r in range(2, est.max_row + 1):
        cod = norm_codigo(est.cell(r, c["codigo"]).value)
        if not cod:
            continue
        filial = (est.cell(r, c["filial"]).value or "").strip() if c["filial"] else ""
        if filial and filial not in vistas:
            vistas.append(filial)
        cor_bruta = texto(est.cell(r, c["cor"]).value)
        codcor_bruto = norm_codigo(est.cell(r, c["codigo_cor"]).value)
        tam_bruto = norm_codigo(est.cell(r, c["tamanho"]).value)
        qtd_bruta = qtd_de(est, "Estoque", r, c["qtde"], "ESTOQUE ATUAL")
        # o banco guarda a linha como o arquivo mandou, inclusive a filial que
        # a planilha vai descartar
        f.linhas["estoque"].append({
            "codigo": cod, "codigo_cor": codcor_bruto, "tamanho": tam_bruto,
            "filial": filial, "qtd": qtd_bruta, "cor": cor_bruta})
        if filiais_estoque and filial not in filiais_estoque:
            continue
        cor = nrm(est.cell(r, c["cor"]).value)
        qtd = qtd_bruta
        f.estoque[cod][cor] += qtd
        f.linhas_estoque.append((
            cod, cor,
            norm_codigo(est.cell(r, c["codigo_cor"]).value),
            norm_codigo(est.cell(r, c["tamanho"]).value),
            qtd,
        ))
    f.filiais_estoque = vistas

    ven = wb["Vendas"]
    c = prepara(ven, COLUNAS_VENDAS, "qtde", "Qtde")
    f.vendas = defaultdict(lambda: defaultdict(lambda: defaultdict(int)))
    filiais_v = []
    for r in range(2, ven.max_row + 1):
        cod = norm_codigo(ven.cell(r, c["codigo"]).value)
        if not cod:
            continue
        filial = (ven.cell(r, c["filial"]).value or "").strip().upper() if c["filial"] else ""
        if filial and filial not in filiais_v:
            filiais_v.append(filial)
        qtd_v = qtd_de(ven, "Vendas", r, c["qtde"], "Qtde")
        f.vendas[cod][nrm(ven.cell(r, c["cor"]).value)][filial] += qtd_v
        f.linhas["vendas"].append({
            "codigo": cod,
            "codigo_cor": norm_codigo(ven.cell(r, c["codigo_cor"]).value) if c["codigo_cor"] else "",
            "tamanho": norm_codigo(ven.cell(r, c["tamanho"]).value) if c["tamanho"] else "",
            "filial": filial, "qtd": qtd_v,
            "cor": texto(ven.cell(r, c["cor"]).value)})
    f.filiais_vendas = filiais_v

    pro = wb["Producao"]
    c = prepara(pro, COLUNAS_PRODUCAO, "qtde", "Quantidade")
    f.producao = defaultdict(lambda: defaultdict(int))
    # A aba mistura peça acabada com insumo consumido: tecido em metros, botão
    # e etiqueta entram com quantidade NEGATIVA e tamanho U. Hoje esses códigos
    # (000076, 000149, 000151) não existem nas abas de trabalho, então não
    # chegam a fazer estrago — mas se um dia colidirem com um código de produto,
    # a produção seria subtraída em silêncio. Produção negativa não existe.
    negativas = 0
    for r in range(2, pro.max_row + 1):
        cod = norm_codigo(pro.cell(r, c["codigo"]).value)
        if not cod:
            continue
        q = qtd_de(pro, "Producao", r, c["qtde"], "Quantidade")
        if q < 0:
            negativas += 1
            continue
        f.producao[cod][nrm(pro.cell(r, c["cor"]).value)] += q
        f.linhas["producao"].append({
            "codigo": cod, "codigo_cor": "",
            "tamanho": norm_codigo(pro.cell(r, c["tamanho"]).value) if c["tamanho"] else "",
            "qtd": q, "cor": texto(pro.cell(r, c["cor"]).value)})
    if negativas:
        f.avisos.append(
            f"Producao: {negativas} linha(s) com quantidade negativa ignorada(s) "
            f"— sao insumo consumido (tecido, botao, etiqueta), nao peca."
        )

    prd = wb["Produtos"]
    c = colunas_da_fonte(prd, COLUNAS_PRODUTOS)
    for r in range(2, prd.max_row + 1):
        f.produtos.append({
            "codigo": norm_codigo(prd.cell(r, c["codigo"]).value),
            "descricao": prd.cell(r, c["descricao"]).value,
            "colecao": prd.cell(r, c["colecao"]).value,
            "divisao": prd.cell(r, c["divisao"]).value,
        })

    for cores in list(f.estoque.values()) + list(f.producao.values()):
        f.cores_conhecidas |= {c for c in cores if c}
    for cores in f.vendas.values():
        f.cores_conhecidas |= {c for c in cores if c}

    wb.close()
    return f


# --------------------------------------------------------------------------- #
# Estrutura das abas Masculino / Feminino
# --------------------------------------------------------------------------- #

CABECALHOS = {
    "estoque_atual": "D",
    "estoque_inicial": "J",
    "consignacao": "E",
    "vendas_totais": "I",
    "sellout": "K",
}


def mapa_colunas(ws) -> dict:
    """Descobre as colunas pelo texto do cabeçalho (linha 3).

    Necessário porque a aba Feminino dos Clássicos tem uma coluna
    "Curadobia" a mais, deslocando tudo à direita dela.
    """
    m = {}
    for c in range(1, 30):
        t = texto(ws.cell(3, c).value)
        if not isinstance(t, str):
            continue
        k = nrm(t)
        if k.startswith("ESTOQUE ATUAL"):
            m["D"] = c
        elif k.startswith("ESTOQUE INICIAL"):
            m["J"] = c
        elif k.startswith("CONSIGNA"):
            m["E"] = c
        elif k == "VENDAS TOTAIS":
            m["I"] = c
        elif k == "VENDAS JARDINS":
            m["JARDINS"] = c
        elif k == "VENDAS IGUATEMI":
            m["IGUATEMI"] = c
        elif k == "VENDAS SITE":
            m["SITE"] = c
        elif k.startswith("SELLOUT"):
            m["K"] = c
    return m


def blocos_de(ws) -> list[dict]:
    """Blocos de coleção: cabeçalho, primeira e última linha de produto, subtotal."""
    blocos = []
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, 2).value != "Código":
            continue
        ini = r + 1
        fim = r
        rr = ini
        while rr <= ws.max_row and ws.cell(rr, 2).value not in (None, "Código"):
            fim = rr
            rr += 1
        blocos.append({
            "hdr": r,
            "titulo": texto(ws.cell(r, 3).value),
            "ini": ini,
            "fim": fim,
            "sub": rr,
            "colecao": colecao_do_bloco(texto(ws.cell(r, 3).value)),
        })
    return blocos


def linhas_de_produto(ws, blocos) -> list[tuple[int, str, dict]]:
    linhas = []
    for b in blocos:
        for r in range(b["ini"], b["fim"] + 1):
            cod = ws.cell(r, 2).value
            if cod in (None, "Código"):
                continue
            linhas.append((r, norm_codigo(cod), b))
    return linhas
