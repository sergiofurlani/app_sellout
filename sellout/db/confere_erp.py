"""A exportação do próprio ERP contra o banco, na mesma semana.

    python -m sellout.db.confere_erp VENDAS_14_20.xlsx --semana 2026-09-20

**Por que isto existe, e por que não é mais uma conferência contra a planilha.**
A planilha não é fonte: é uma consumidora anterior da mesma fonte. Quando ela e
o banco discordam, comparar um com o outro não diz quem está certo — diz apenas
que discordam. A autoridade é o ERP.

A exportação do ERP vem no mesmo grão da tabela `venda` (filial, código, cor,
tamanho, quantidade, devolução com sinal negativo) e é produzida por um caminho
que não é o nosso: relatório do sistema, não a nossa extração pela API. Então
ela responde a pergunta que importa antes de aposentar a planilha:

    a nossa extração é fiel ao ERP naquela semana?

Se for, o que sobra de divergência contra a planilha é história da planilha — e
aí o arquivo pode ser aposentado em vez de reconciliado. Se não for, achamos um
defeito na extração, que é o único dos três lugares onde um defeito ainda
importa.

**A tolerância é zero.** Mesma semana, mesma fonte, mesmo grão: qualquer peça
de diferença é achado, não arredondamento. É o oposto da comparação contra a
planilha, onde uma peça sai de data de corte.

Um detalhe que esta conferência pega de graça: duas cargas da mesma semana
dobram tudo, e o número dobrado continua parecendo plausível. Por isso o
programa conta os snapshots antes de somar e para se houver mais de um.
"""

from __future__ import annotations

import argparse
import pathlib
from collections import defaultdict
from datetime import date

import openpyxl

from ..core.leitura import norm_codigo
from .conexao import conectar

# O cabeçalho da exportação. Achado pelo texto, não pela posição: relatório de
# ERP ganha coluna sem avisar, e foi assim que a coluna "Consignado Real"
# apareceu no meio da planilha.
COLUNAS = {"filial": ("FILIAL",),
           "codigo": ("CÓDIGO", "CODIGO"),
           "codigo_cor": ("CODIGO_COR", "CÓDIGO_COR", "COD_COR", "CODIGO COR"),
           "tamanho": ("TAM", "TAMANHO"),
           "qtd": ("QTDE", "QTD", "QUANTIDADE")}


# **A mesma loja com dois nomes.** O relatório do ERP escreve JARDINS; a API
# devolve o código da filial, EGREY JDS. Não é defeito de dado nem de carga — e
# sem isto a conferência acusa a loja inteira dos dois lados, 107 chaves "só na
# exportação" e 104 "só no banco", como se nada batesse. Foi o que aconteceu na
# primeira corrida, em 27/09.
ALIAS_FILIAL = {"JARDINS": "EGREY JDS"}


def nrm(v) -> str:
    return str(v or "").strip().upper()


def filial_do_export(nome) -> str:
    """O nome do relatório traduzido para o código da filial do banco."""
    n = nrm(nome)
    return ALIAS_FILIAL.get(n, n)


def so_de_um_lado(export: dict, banco: dict) -> dict:
    """{export: [filiais], banco: [filiais]} — nome que existe de um lado só.

    É o sintoma de apelido faltando, e tem de gritar: uma loja que não casa por
    nome produz centenas de divergências falsas e nenhuma pista de que o
    problema é ortográfico. Vale para a loja nova que ninguém mapeou ainda.
    """
    fe = {k[2] for k in export}
    fb = {k[2] for k in banco}
    return {"export": sorted(fe - fb), "banco": sorted(fb - fe)}


def mapa_do_export(ws) -> dict:
    """{campo: coluna} pela linha de cabeçalho."""
    m = {}
    for c in range(1, ws.max_column + 1):
        t = nrm(ws.cell(1, c).value)
        for campo, nomes in COLUNAS.items():
            if t in nomes and campo not in m:
                m[campo] = c
    return m


def do_export(caminho: str) -> tuple[dict, list]:
    """{(codigo, codigo_cor, filial): qtd} — os tamanhos somados.

    O tamanho é somado porque a planilha e o sellout são por produto e cor; e
    porque uma troca de tamanho dentro do mesmo produto não muda nada do que
    está sendo medido aqui.
    """
    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = wb[wb.sheetnames[0]]
    m = mapa_do_export(ws)
    faltando = [c for c in COLUNAS if c not in m and c != "tamanho"]
    if faltando:
        wb.close()
        raise ValueError(f"a exportacao nao tem as colunas {faltando}; "
                         f"cabecalho lido: {[ws.cell(1, c).value for c in range(1, ws.max_column + 1)]}")
    fora, avisos = defaultdict(float), []
    for r in range(2, ws.max_row + 1):
        cod = norm_codigo(ws.cell(r, m["codigo"]).value)
        if not cod:
            continue
        q = ws.cell(r, m["qtd"]).value
        if not isinstance(q, (int, float)):
            avisos.append(f"linha {r}: quantidade nao numerica ({q!r})")
            continue
        chave = (cod, nrm(ws.cell(r, m["codigo_cor"]).value),
                 filial_do_export(ws.cell(r, m["filial"]).value))
        fora[chave] += float(q)
    wb.close()
    return dict(fora), avisos


def do_csv(caminho: str) -> tuple[dict, list]:
    """O mesmo dicionário, lido do CSV que o `coletor.vendas` grava.

    **É o que separa dois defeitos que dão o mesmo sintoma.** Se a exportação do
    ERP tem uma devolução que o banco não tem, ela se perdeu na extração ou na
    carga — e o CSV do coletor fica exatamente no meio dos dois. Comparar com
    ele diz qual metade olhar, em vez de ler o código das duas.

    O coletor também descarta chave de saldo zero (`abs(qtd) > 0`), igual à
    carga: uma venda e uma devolução da mesma peça na mesma semana não viram
    linha em lugar nenhum.
    """
    import csv as _csv
    with open(caminho, newline="", encoding="utf-8-sig") as f:
        linhas = list(_csv.reader(f, delimiter=";"))
    if not linhas:
        raise ValueError(f"{caminho} esta vazio")
    cab = [nrm(x) for x in linhas[0]]
    m = {}
    for i, t in enumerate(cab, 1):
        for campo, nomes in COLUNAS.items():
            if t in nomes and campo not in m:
                m[campo] = i
    faltando = [c for c in COLUNAS if c not in m and c != "tamanho"]
    if faltando:
        raise ValueError(f"o CSV nao tem as colunas {faltando}; cabecalho: {cab}")
    fora, avisos = defaultdict(float), []
    for n, linha in enumerate(linhas[1:], 2):
        if len(linha) < max(m.values()):
            continue
        cod = norm_codigo(linha[m["codigo"] - 1])
        if not cod:
            continue
        try:
            q = float(str(linha[m["qtd"] - 1]).replace(",", "."))
        except ValueError:
            avisos.append(f"linha {n}: quantidade nao numerica "
                          f"({linha[m['qtd'] - 1]!r})")
            continue
        fora[(cod, nrm(linha[m["codigo_cor"] - 1]),
              filial_do_export(linha[m["filial"] - 1]))] += q
    return dict(fora), avisos


def snapshots_da_semana(semana: date, origem: str) -> list:
    with conectar() as c, c.cursor() as cur:
        cur.execute("SELECT id FROM snapshot WHERE data = %s AND origem = %s "
                    "ORDER BY id", (semana, origem))
        return [r[0] for r in cur.fetchall()]


def do_banco(snaps) -> dict:
    """{(codigo, codigo_cor, filial): qtd} para os snapshots dados."""
    if not snaps:
        return {}
    fora = defaultdict(float)
    with conectar() as c, c.cursor() as cur:
        cur.execute(
            "SELECT v.codigo, v.codigo_cor, v.filial, sum(v.qtd) "
            "  FROM venda v WHERE v.snapshot_id = ANY(%s) "
            " GROUP BY v.codigo, v.codigo_cor, v.filial", (list(snaps),))
        for cod, cor, fil, q in cur.fetchall():
            fora[(norm_codigo(cod), nrm(cor), nrm(fil))] += float(q)
    return dict(fora)


def compara(export: dict, banco: dict) -> dict:
    """Chave a chave, com tolerância zero.

    Chave que existe de um lado só **não** entra como diferença de número: é
    linha que um dos dois não viu, e somar zero do outro lado misturaria as duas
    coisas — foi o erro que a conferência contra a planilha já tinha corrigido.
    """
    iguais, difere, so_erp, so_banco = [], [], [], []
    for k in sorted(set(export) | set(banco)):
        e, b = export.get(k), banco.get(k)
        if e is None:
            so_banco.append({"chave": k, "qtd": b})
        elif b is None:
            so_erp.append({"chave": k, "qtd": e})
        elif abs(e - b) < 1e-9:
            iguais.append(k)
        else:
            difere.append({"chave": k, "erp": e, "banco": b, "dif": e - b})
    return {"iguais": iguais, "difere": difere,
            "so_erp": so_erp, "so_banco": so_banco,
            "total_erp": sum(export.values()),
            "total_banco": sum(banco.values())}


def por_filial(r: dict) -> dict:
    """{filial: {erp, banco, dif, chaves}} — onde a diferença mora.

    Agrupa as três listas juntas: uma filial que só existe num dos lados é
    exatamente o caso que soma errado em silêncio, e ela tem de aparecer na
    mesma tabela que as outras.
    """
    fora = defaultdict(lambda: {"erp": 0.0, "banco": 0.0, "dif": 0.0, "chaves": 0})
    def soma(filial, e, b):
        d = fora[filial]
        d["erp"] += e
        d["banco"] += b
        d["dif"] += e - b
        if abs(e - b) > 1e-9:
            d["chaves"] += 1
    for x in r["difere"]:
        soma(x["chave"][2], x["erp"], x["banco"])
    for x in r["so_erp"]:
        soma(x["chave"][2], x["qtd"], 0.0)
    for x in r["so_banco"]:
        soma(x["chave"][2], 0.0, x["qtd"])
    return {f: dict(v) for f, v in fora.items()}


def fiel(r: dict) -> bool:
    """A extração reproduz o ERP naquela semana, chave a chave."""
    return not r["difere"] and not r["so_erp"] and not r["so_banco"]


def lista_snapshots(limite=15) -> list:
    """As últimas semanas gravadas, com origem e linhas de venda.

    Existe porque "o banco nao tem snapshot dessa data" não diz qual data tem, e
    a rodada do app grava no dia em que roda, não no domingo que fecha a semana.
    """
    with conectar() as c, c.cursor() as cur:
        cur.execute(
            "SELECT s.data, s.origem, s.quem, count(v.snapshot_id) AS linhas "
            "  FROM snapshot s LEFT JOIN venda v ON v.snapshot_id = s.id "
            " GROUP BY s.id, s.data, s.origem, s.quem "
            " ORDER BY s.data DESC, s.id DESC LIMIT %s", (limite,))
        return [{"data": d, "origem": o, "quem": q or "", "linhas": n}
                for d, o, q, n in cur.fetchall()]


def main(argv=None):
    p = argparse.ArgumentParser(
        description="A exportacao do ERP contra o banco, na mesma semana")
    p.add_argument("export", nargs="?",
                   help="xlsx exportado do ERP (Filial, Codigo, Qtde...)")
    p.add_argument("--semana", type=date.fromisoformat,
                   help="o domingo que fecha a semana, como esta em snapshot.data")
    p.add_argument("--origem", default="erp", choices=("erp", "upload"),
                   help="erp = carga retroativa; upload = rodada do app")
    p.add_argument("--csv", help="compara com o CSV do coletor.vendas em vez\n"
                                "                         do banco: separa extracao de carga")
    p.add_argument("--listar", action="store_true",
                   help="so lista as ultimas semanas gravadas e sai")
    p.add_argument("--limite", type=int, default=20)
    args = p.parse_args(argv)

    if args.listar:
        print("\nUltimas semanas no banco:\n")
        print(f'  {"data":12} {"origem":8} {"quem":22} {"linhas":>8}')
        print("  " + "-" * 54)
        for s_ in lista_snapshots():
            print(f'  {str(s_["data"]):12} {s_["origem"]:8} {s_["quem"][:22]:22} '
                  f'{s_["linhas"]:>8,}')
        return 0

    if not args.export:
        p.error("falta o arquivo exportado do ERP")
    if not args.semana and not args.csv:
        p.error("--semana e obrigatorio (ou use --csv)")

    caminho = pathlib.Path(args.export.strip().strip("<>").strip('"').strip("'"))
    if not caminho.exists():
        print(f"\nNao encontrei o arquivo: {caminho}")
        return 1

    ex, avisos = do_export(str(caminho))
    rotulo = "extracao" if args.csv else "banco"
    if args.csv:
        print(f"\nERP x extracao (CSV do coletor) — {args.csv}")
    else:
        print(f"\nERP x banco — semana de {args.semana}, origem {args.origem}")
    print(f"  exportacao: {len(ex):,} chave(s) produto+cor+filial")
    if avisos:
        print(f"  {len(avisos)} linha(s) ignorada(s) na exportacao:")
        for a in avisos[:5]:
            print(f"    {a}")

    if args.csv:
        ba, avisos_csv = do_csv(args.csv)
        for a in avisos_csv[:5]:
            print(f"  CSV: {a}")
        snaps = None
    else:
        snaps = snapshots_da_semana(args.semana, args.origem)
    if snaps is not None and not snaps:
        print(f"\n  O banco nao tem snapshot de {args.semana} com origem "
              f"'{args.origem}'.")
        print("  Sem isso nao ha o que comparar. `--listar` diz quais datas tem:")
        print("  a rodada do app grava no dia em que roda, nao no domingo que")
        print("  fecha a semana.")
        return 1
    if snaps is not None and len(snaps) > 1:
        # Duas cargas da mesma semana dobram tudo, e o dobro parece plausivel.
        print(f"\n  {len(snaps)} snapshots para a MESMA semana e origem: {snaps}")
        print("  Somar os dois dobraria a venda da semana. Isto e defeito de")
        print("  carga, e precisa ser resolvido antes de conferir qualquer coisa.")
        return 1

    if snaps is not None:
        ba = do_banco(snaps)
    r = compara(ex, ba)
    print(f"  {rotulo + ':':11} {len(ba):,} chave(s)")

    # Antes de qualquer número: nome que existe de um lado só invalida a leitura
    # de tudo o que vem depois.
    sozinhas = so_de_um_lado(ex, ba)
    if sozinhas["export"] or sozinhas["banco"]:
        print("\n  FILIAL COM NOME DE UM LADO SO — leia o resto com desconfianca:")
        if sozinhas["export"]:
            print(f"    so na exportacao do ERP: {', '.join(sozinhas['export'])}")
        if sozinhas["banco"]:
            print(f"    so no banco:             {', '.join(sozinhas['banco'])}")
        print("    Se for a mesma loja com dois nomes, a conferencia acusa a loja")
        print("    inteira duas vezes e nada bate. O apelido vai em ALIAS_FILIAL.")
    print(f"\n  total de pecas   ERP {r['total_erp']:>8,.0f}   "
          f"{rotulo} {r['total_banco']:>8,.0f}   "
          f"dif {r['total_erp'] - r['total_banco']:>+7,.0f}")
    print(f"\n  iguais na peca          {len(r['iguais']):>6}")
    print(f"  diferem                 {len(r['difere']):>6}")
    print(f"  so na exportacao do ERP {len(r['so_erp']):>6}")
    print(f"  so no {rotulo:18} {len(r['so_banco']):>6}")

    if fiel(r):
        print("\n  A EXTRACAO REPRODUZ O ERP NESTA SEMANA, chave a chave.")
        print("  Então o que sobra de divergencia contra a planilha e historia da")
        print("  planilha, e o arquivo pode ser aposentado em vez de reconciliado.")
        print("  Uma semana nao prova as 38 — repita nas semanas que tiverem")
        print("  exportacao guardada.")
        return 0

    pf = por_filial(r)
    print("\n  onde a diferenca mora:")
    print(f'    {"filial":14} {"erp":>8} {"banco":>8} {"dif":>7} {"chaves":>7}')
    print("    " + "-" * 50)
    for f, d in sorted(pf.items(), key=lambda kv: -abs(kv[1]["dif"])):
        print(f'    {f[:14]:14} {d["erp"]:>8,.0f} {d["banco"]:>8,.0f} '
              f'{d["dif"]:>+7,.0f} {d["chaves"]:>7}')

    for rotulo, lista in (("diferem", r["difere"]), ("so na exportacao", r["so_erp"]),
                          (f"so no {rotulo}", r["so_banco"])):
        if not lista:
            continue
        print(f"\n  {rotulo} ({len(lista)}):")
        print(f'    {"codigo":8} {"cor":6} {"filial":14} {"erp":>7} {"banco":>7}')
        print("    " + "-" * 48)
        for x in lista[:args.limite]:
            cod, cor, fil = x["chave"]
            e = x.get("erp", x.get("qtd") if lista is r["so_erp"] else 0.0)
            b = x.get("banco", x.get("qtd") if lista is r["so_banco"] else 0.0)
            print(f'    {cod:8} {cor:6} {fil[:14]:14} {e:>7,.0f} {b:>7,.0f}')
    print("\n  Tolerancia zero: mesma semana, mesma fonte, mesmo grao. Uma peca")
    print("  de diferenca aqui e achado, nao arredondamento.")
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
