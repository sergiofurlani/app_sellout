"""Carrega o histórico de sellout que mora nas colunas das abas de trabalho.

    python -m sellout.db.historico "sellout geral.xlsx" --planilha geral
    python -m sellout.db.historico "Sellout Clássicos.xlsx" --planilha classicos
    python -m sellout.db.historico geral.xlsx --so-ler     # não grava, só conta

Cada coluna à direita da K é uma rodada: a data no cabeçalho da linha 3 e o
percentual na célula. São 171 semanas no Masculino, desde 04/12/2022.

**O que se recupera é só o resultado.** Vendas, estoque e estoque inicial
daquela semana não existem mais — a planilha sobrescrevia a cada rodada. Por
isso vai para `sellout_historico` e não para `snapshot`: o percentual é
verdade, o que o produziu se perdeu.

Rodar duas vezes não duplica: a chave é (planilha, aba, código, cores, data).
"""

from __future__ import annotations

import argparse
import datetime
from collections import Counter

import openpyxl

from ..core.leitura import ABAS_TRABALHO, blocos_de, linhas_de_produto, num, texto, vermelho
from .conexao import conectar, por_que_nao

VAZIOS = {"-", "--", "", "#DIV/0!", "#N/D", "#N/A", "#REF!", "#VALUE!"}


def colunas_de_data(ws):
    """As colunas de rodada, pela data do cabeçalho. Devolve (colunas, avisos).

    Só data serve como marca: o rótulo da coluna atual é texto ("Sellout
    31/08") e as colunas mais antigas perderam o cabeçalho. Sem data, a coluna
    não entra — chutar a semana de um percentual seria inventar histórico.

    As colunas andam para trás no tempo. Uma data que não é anterior à última
    aceita está errada, e um percentual arquivado na semana errada é pior que
    um ausente, porque ninguém vai conferir de novo.

    **Quando tirar um ano conserta a ordem, conserta.** O defeito é sempre o
    mesmo e aparece todo fim de ano: em janeiro alguém digita o ano novo numa
    coluna de dezembro. `2026-12-15` entre `2026-01-04` e `2025-12-07` era
    2025; o mesmo acontece em 2025-12-29, 2024-12-17 e 2023-12-11. Descartar
    essas semanas jogaria fora histórico bom por um erro de digitação
    reconhecível.
    """
    brutas = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(3, c).value
        if isinstance(v, datetime.datetime):
            brutas.append((c, v.date()))
        elif isinstance(v, datetime.date):
            brutas.append((c, v))

    boas, avisos = [], []
    ultima = None
    for c, d in brutas:
        if ultima is None or d < ultima:
            boas.append((c, d))
            ultima = d
            continue
        try:
            um_ano_antes = d.replace(year=d.year - 1)
        except ValueError:                         # 29 de fevereiro
            um_ano_antes = d.replace(year=d.year - 1, day=28)
        if um_ano_antes < ultima:
            avisos.append(f"coluna {c}: {d} corrigida para {um_ano_antes} "
                          "(ano digitado errado; a ordem volta a fechar)")
            boas.append((c, um_ano_antes))
            ultima = um_ano_antes
        else:
            avisos.append(f"coluna {c}: {d} nao e anterior a {ultima} e nao da "
                          "para corrigir; semana descartada")
    return boas, avisos


def percentual(valor):
    """Devolve a fração, ou None quando a célula não diz nada.

    `-` é o que a planilha usa para "não havia produto nessa semana": 2.333
    células. Tratar como zero afundaria a média de qualquer consulta.
    """
    if valor is None:
        return None
    if isinstance(valor, str) and valor.strip() in VAZIOS:
        return None
    n = num(valor)
    if n is None:
        return None
    # a coluna é fração (0,2063 = 20,63%); se alguém digitou 20,63, converte
    return n / 100 if n > 1.5 else n


def ler(caminho: str, planilha: str) -> tuple[list[dict], dict, list[str]]:
    wb = openpyxl.load_workbook(caminho, data_only=True)
    wr = openpyxl.load_workbook(caminho, rich_text=True)
    registros = []
    resumo = Counter()
    avisos: list[str] = []
    datas = set()

    for aba in ABAS_TRABALHO:
        if aba not in wb.sheetnames:
            continue
        ws, wsr = wb[aba], wr[aba]
        colunas, suspeitas = colunas_de_data(ws)
        resumo[f"{aba}: colunas de data"] = len(colunas)
        avisos.extend(f"{aba} — {s}" for s in suspeitas)
        if not colunas:
            continue
        datas.update(d for _c, d in colunas)

        todas = [(l, c, b) for l, c, b in linhas_de_produto(ws, blocos_de(ws)) if c]
        # Duas linhas com o mesmo código e o mesmo texto em vermelho não têm
        # como ser distinguidas — é o caso do 330011 (D2), que aparece duas
        # vezes no Feminino sem vermelho nenhum, com percentuais DIFERENTES em
        # 23 semanas. Gravar uma delas seria escolher no escuro, então nenhuma
        # entra e o aviso diz qual ficou de fora.
        repetidas = Counter((c, vermelho(wsr.cell(l, 3).value)) for l, c, _b in todas)
        ambiguas = {k for k, n in repetidas.items() if n > 1}
        for codigo, cores in sorted(ambiguas):
            avisos.append(
                f"{aba} — codigo {codigo} aparece em mais de uma linha com o "
                f"mesmo vermelho ({cores or 'nenhum'}); historico nao carregado")

        for linha, codigo, bloco in todas:
            cores = vermelho(wsr.cell(linha, 3).value)
            if (codigo, cores) in ambiguas:
                resumo["linhas ambiguas puladas"] += 1
                continue
            descricao = texto(ws.cell(linha, 3).value)
            for coluna, data in colunas:
                p = percentual(ws.cell(linha, coluna).value)
                if p is None:
                    resumo["celulas vazias"] += 1
                    continue
                registros.append({
                    "planilha": planilha, "aba": aba, "codigo": codigo,
                    "cores": cores, "data": data, "percentual": p,
                    "bloco": bloco.get("titulo"), "colecao": bloco.get("colecao"),
                    "descricao": descricao,
                })
    resumo["registros"] = len(registros)
    resumo["semanas"] = len(datas)
    if datas:
        resumo["de"] = min(datas)
        resumo["ate"] = max(datas)
    return registros, dict(resumo), avisos


def gravar(registros: list[dict]) -> int:
    """Grava, sobrescrevendo o mesmo (planilha, aba, código, cores, data)."""
    if not registros:
        return 0
    sql = (
        "INSERT INTO sellout_historico "
        "(planilha, aba, codigo, cores, data, percentual, bloco, colecao, descricao) "
        "VALUES (%(planilha)s, %(aba)s, %(codigo)s, %(cores)s, %(data)s, "
        "        %(percentual)s, %(bloco)s, %(colecao)s, %(descricao)s) "
        "ON CONFLICT (planilha, aba, codigo, cores, data) DO UPDATE SET "
        "  percentual = EXCLUDED.percentual, bloco = EXCLUDED.bloco, "
        "  colecao = EXCLUDED.colecao, descricao = EXCLUDED.descricao, "
        "  carregado_em = now()")
    with conectar() as c, c.cursor() as cur:
        cur.executemany(sql, registros)
    return len(registros)


def main(argv=None):
    p = argparse.ArgumentParser(description="Carrega o historico de sellout no banco")
    p.add_argument("arquivo")
    p.add_argument("--planilha", choices=("geral", "classicos"), required=True)
    p.add_argument("--so-ler", action="store_true", help="nao grava, so conta")
    args = p.parse_args(argv)

    registros, resumo, avisos = ler(args.arquivo, args.planilha)
    print(f"\n{args.arquivo}  ({args.planilha})")
    for k, v in resumo.items():
        print(f"  {k:28} {v}")
    for a in avisos:
        print(f"  AVISO: {a}")

    if args.so_ler:
        return 0
    motivo = por_que_nao()
    if motivo:
        print(f"\n  Nao gravei: {motivo}")
        return 1
    print(f"\n  gravados: {gravar(registros)}")
    with conectar() as c, c.cursor() as cur:
        cur.execute("SELECT count(*), count(DISTINCT data), min(data), max(data) "
                    "FROM sellout_historico")
        n, semanas, de, ate = cur.fetchone()
        print(f"  no banco: {n} registros, {semanas} semanas, de {de} a {ate}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
