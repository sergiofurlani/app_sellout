"""Lê o estoque e o preço exportados do ERP e grava como retrato da semana.

    python -m coletor.fontes_erp --estoque ESTOQUE.xls --preco Tabela_preco.xlsx
    python -m coletor.fontes_erp --estoque ESTOQUE.xls --preco Tabela_preco.xlsx --aplicar

São as duas fontes que a API não entregou. O `estoque/Lista` e o
`MILLENIUM_ECO.PRODUTOS.SALDODEESTOQUE` existem e respondem, mas nenhum dos
dois devolveu linha útil sem um filtro que ninguém soube dizer qual é — e
insistir já custou caro uma vez. Enquanto isso, o ERP exporta os dois, e o
formato é estável.

**Isto não é a planilha de trabalho.** São dois arquivos crus, direto do ERP,
com as mesmas colunas das abas Estoque e Preco. O que muda é que eles vão para
o banco sem passar pela `sellout geral.xlsx` — a planilha deixa de ser o
caminho do dado e vira só saída.

O estoque chega nos dois formatos que o ERP exporta: um `.xls` que por dentro
é **SpreadsheetML** (XML da Microsoft, que openpyxl não abre) e um `.xlsx` de
verdade. O leitor é escolhido pelos primeiros bytes do arquivo, não pela
extensão — confiar na extensão erra justamente no primeiro.

Quando a API de integração destravar, troca-se `le_estoque` e `le_preco` por
chamadas e **nada mais muda**: o resto do caminho já está pronto.
"""

from __future__ import annotations

import argparse
import pathlib
import xml.etree.ElementTree as ET
from collections import Counter
from datetime import date

NS = {"ss": "urn:schemas-microsoft-com:office:spreadsheet"}

# Filial do ERP -> como o resto do projeto chama. O e-commerce (00044) não
# aparece neste export, e é assim que deve ser: ele vende do estoque das lojas
# e contá-lo dobraria o saldo (D12).
FILIAIS = {
    "EGREY IGUATEMI": "IGUATEMI",
    "EGREY JARDINS": "EGREY JDS",
    "EGREY JDS": "EGREY JDS",
    "IGUATEMI": "IGUATEMI",
}

# Códigos que o ERP guarda junto com a roupa mas não são peça de venda.
# Retirado a pedido do negócio em 21/09: o `SW4-WH-FSC` sozinho carregava
# 1.000 das 10.083 peças — 10% do estoque — e não tem preço, porque não é
# vendido. Entrava inflando o denominador do sellout de todo mundo.
#
# A lista é explícita de propósito. Adivinhar pelo formato do código, pelo
# tamanho `U` ou pela falta de preço já deu errado neste projeto mais de uma
# vez; quem decide o que é peça é o negócio, não o padrão do texto.
FORA_DO_SELLOUT = {"SW4-WH-FSC"}


def dia(t: str) -> date:
    return date.fromisoformat(t)


def _texto(celula) -> str:
    d = celula.find("ss:Data", NS)
    return (d.text or "").strip() if d is not None else ""


def _numero(t: str) -> float:
    t = (t or "").strip().replace(".", "").replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return 0.0


def _do_spreadsheetml(caminho: str) -> list[list[str]]:
    """`.xls` que por dentro é XML da Microsoft. openpyxl não abre."""
    raiz = ET.parse(caminho).getroot()
    tabela = raiz.find(".//ss:Table", NS)
    if tabela is None:
        return []
    return [[_texto(c) for c in r.findall("ss:Cell", NS)]
            for r in tabela.findall("ss:Row", NS)]


def _do_xlsx(caminho: str) -> list[list[str]]:
    import openpyxl
    wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    return [["" if v is None else str(v).strip() for v in linha]
            for linha in ws.iter_rows(values_only=True)]


def _linhas(caminho: str) -> list[list[str]]:
    """Escolhe o leitor pelo conteúdo, não pela extensão.

    O ERP exporta os dois: um `.xls` que é XML e um `.xlsx` de verdade. Confiar
    na extensão erra justamente no primeiro, que é o que ele manda por padrão.
    """
    with open(caminho, "rb") as f:
        inicio = f.read(4)
    if inicio[:2] == b"PK":
        return _do_xlsx(caminho)
    return _do_spreadsheetml(caminho)


def le_estoque(caminho: str) -> tuple[list[dict], dict, list[str]]:
    """Export de estoque -> linhas por produto, cor, tamanho e filial."""
    linhas = _linhas(caminho)
    if not linhas:
        return [], {}, [f"{caminho}: planilha vazia"]

    cab = [c.strip().upper() for c in linhas[0]]
    def onde(*nomes):
        for n in nomes:
            if n in cab:
                return cab.index(n)
        return None

    i_filial = onde("FILIAL")
    i_codigo = onde("CÓDIGO", "CODIGO PRODUTO", "COD_PRODUTO")
    i_desc = onde("DESCRIÇÃO", "DESCRICAO")
    i_cor = onde("CODIGO_COR", "CODIGO COR", "COD_COR")
    # Na primeira versão do export a coluna da cor se chamava `CODIGO`, igual à
    # do produto depois de tirar o acento. A partir de 21/09 ela vem como
    # `CODIGO_COR`; o desempate por posição fica para quem ainda tiver um
    # arquivo antigo na mão.
    if i_cor is None and i_desc is not None:
        # o cabeçalho repete "CODIGO"; a segunda ocorrência é a da cor
        for j in range(i_desc + 1, len(cab)):
            if cab[j].startswith("CODIGO"):
                i_cor = j
                break
    i_nome_cor = onde("COR")
    i_tam = onde("TAMANHO")
    i_qtd = onde("ESTOQUE ATUAL", "QUANTIDADE", "SALDO")

    faltando = [n for n, i in (("Filial", i_filial), ("Código", i_codigo),
                               ("código da cor", i_cor), ("tamanho", i_tam),
                               ("Estoque atual", i_qtd)) if i is None]
    if faltando:
        return [], {}, [f"{caminho}: faltam colunas: {', '.join(faltando)}"]

    saida, avisos = [], []
    resumo = Counter()
    desconhecidas = set()
    for l in linhas[1:]:
        if len(l) <= max(i_filial, i_codigo, i_cor, i_tam, i_qtd):
            continue
        bruta = l[i_filial].strip()
        filial = FILIAIS.get(bruta.upper())
        if filial is None:
            desconhecidas.add(bruta)
            filial = bruta
        codigo = l[i_codigo].strip()
        if not codigo:
            continue
        if codigo.upper() in FORA_DO_SELLOUT:
            resumo["fora do sellout"] += 1
            resumo["pecas fora"] += _numero(l[i_qtd])
            continue
        qtd = _numero(l[i_qtd])
        resumo["linhas"] += 1
        resumo["pecas"] += qtd
        saida.append({
            "codigo": codigo,
            "codigo_cor": l[i_cor].strip(),
            "cor": l[i_nome_cor].strip() if i_nome_cor is not None else "",
            "tamanho": l[i_tam].strip(),
            "filial": filial,
            "qtd": qtd,
        })
    if desconhecidas:
        avisos.append("filial fora do mapa FILIAIS, gravada como veio: "
                      + ", ".join(sorted(desconhecidas)))
    resumo["filiais"] = len({r["filial"] for r in saida})
    resumo["produtos"] = len({r["codigo"] for r in saida})
    return saida, dict(resumo), avisos


def le_preco(caminho: str) -> tuple[list[dict], dict, list[str]]:
    """xlsx do ERP -> preço por produto, cor e tamanho."""
    import openpyxl

    wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    cab = [str(c or "").strip().upper() for c in next(it)]

    def onde(*nomes):
        for n in nomes:
            if n in cab:
                return cab.index(n)
        return None

    i_prod, i_cor = onde("PRODUTO", "CÓDIGO", "COD_PRODUTO"), onde("COR", "COD_COR")
    i_tam, i_preco = onde("TAMANHO"), onde("PRECO", "PREÇO")
    faltando = [n for n, i in (("Produto", i_prod), ("Cor", i_cor),
                               ("Tamanho", i_tam), ("Preco", i_preco)) if i is None]
    if faltando:
        return [], {}, [f"{caminho}: faltam colunas: {', '.join(faltando)}"]

    saida, avisos = [], []
    resumo = Counter()
    for r in it:
        if not r or r[i_prod] in (None, ""):
            continue
        preco = r[i_preco]
        if preco in (None, ""):
            resumo["sem preco"] += 1
            continue
        saida.append({
            "codigo": str(r[i_prod]).strip(),
            "codigo_cor": str(r[i_cor] or "").strip(),
            "tamanho": str(r[i_tam] or "").strip(),
            "preco": float(preco),
        })
        resumo["linhas"] += 1
    resumo["produtos"] = len({r["codigo"] for r in saida})
    return saida, dict(resumo), avisos


def confere(estoque, preco) -> dict:
    """O que o estoque tem e o preço não cobre. É o furo que o Nível de
    Estoque (D7) esconderia caindo no preço de outro tamanho."""
    com_preco = {(p["codigo"], p["codigo_cor"], p["tamanho"]) for p in preco}
    chaves = {(e["codigo"], e["codigo_cor"], e["tamanho"]) for e in estoque}
    orfas = chaves - com_preco
    por_chave = {}
    for e in estoque:
        k = (e["codigo"], e["codigo_cor"], e["tamanho"])
        if k in orfas:
            por_chave[k] = por_chave.get(k, 0.0) + e["qtd"]
    return {"chaves": len(chaves), "sem_preco": len(orfas),
            "pecas_sem_preco": sum(por_chave.values()),
            "exemplos": sorted(orfas)[:8],
            "sem_preco_com_saldo": sorted(por_chave.items(), key=lambda x: -x[1])}


def grava(estoque, preco, data_base: date, observacao: str = "") -> int:
    from sellout.db.conexao import conectar

    with conectar() as c, c.cursor() as cur:
        cur.execute(
            "INSERT INTO snapshot (data, origem, quem, observacao) "
            "VALUES (%s, 'erp', %s, %s) RETURNING id",
            (data_base, "coletor.fontes_erp", observacao or None))
        sid = cur.fetchone()[0]
        cur.executemany(
            "INSERT INTO estoque (snapshot_id, codigo, codigo_cor, tamanho, filial, qtd) "
            "VALUES (%(sid)s, %(codigo)s, %(codigo_cor)s, %(tamanho)s, %(filial)s, %(qtd)s) "
            "ON CONFLICT DO NOTHING",
            [dict(e, sid=sid) for e in estoque])
        cur.executemany(
            "INSERT INTO preco (snapshot_id, codigo, codigo_cor, tamanho, preco) "
            "VALUES (%(sid)s, %(codigo)s, %(codigo_cor)s, %(tamanho)s, %(preco)s) "
            "ON CONFLICT DO NOTHING",
            [dict(p, sid=sid) for p in preco])
    return sid


def main(argv=None):
    p = argparse.ArgumentParser(
        description="Le o estoque e o preco exportados do ERP")
    p.add_argument("--estoque", required=True)
    p.add_argument("--preco", required=True)
    p.add_argument("--data", type=dia, default=date.today(),
                   help="data do retrato; padrao hoje")
    p.add_argument("--aplicar", action="store_true", help="sem isto, so mostra")
    args = p.parse_args(argv)

    for rotulo, caminho in (("estoque", args.estoque), ("preco", args.preco)):
        if not pathlib.Path(caminho).exists():
            print(f"\nNao encontrei o {rotulo}: {caminho}")
            return 1

    est, r_est, av_est = le_estoque(args.estoque)
    pre, r_pre, av_pre = le_preco(args.preco)

    print(f"\nestoque  {args.estoque}")
    for k, v in sorted(r_est.items()):
        print(f"  {k:22} {v:,.0f}" if isinstance(v, float) else f"  {k:22} {v}")
    for a in av_est:
        print(f"  AVISO: {a}")
    print(f"\npreco    {args.preco}")
    for k, v in sorted(r_pre.items()):
        print(f"  {k:22} {v}")
    for a in av_pre:
        print(f"  AVISO: {a}")

    if not est or not pre:
        print("\n  Faltou uma das duas fontes. Nada a gravar.")
        return 1

    c = confere(est, pre)
    print(f"\n  {c['chaves']:,} combinacao(oes) produto+cor+tamanho no estoque")
    print(f"  {c['sem_preco']} sem preco ({c['pecas_sem_preco']:,.0f} peca(s))")
    if c["exemplos"]:
        print("  exemplos:", ", ".join("/".join(x) for x in c["exemplos"]))
    # Sem preço E com saldo é o perfil de insumo, como o SW4-WH-FSC que saiu.
    # Quem decide é o negócio; aqui só aparece quem se parece.
    suspeitos = [(k, q) for k, q in c["sem_preco_com_saldo"] if q >= 10]
    if suspeitos:
        print("\n  Sem preco e com saldo — confira se e peca de venda mesmo:")
        for (cod, cc, tam), q in suspeitos:
            print(f"    {cod}/{cc}/{tam:<4} {q:>8,.0f}   "
                  "(se nao for, acrescente em FORA_DO_SELLOUT)")

    por_filial = Counter()
    for e in est:
        por_filial[e["filial"]] += e["qtd"]
    print("\n  estoque por filial:")
    for f, q in sorted(por_filial.items(), key=lambda x: -x[1]):
        print(f"    {f:14} {q:>9,.0f}")

    if not args.aplicar:
        print("\n  Ensaio. Repita com --aplicar para gravar no banco.")
        return 0

    from sellout.db.conexao import por_que_nao
    motivo = por_que_nao()
    if motivo:
        print(f"\n  Nao gravei: {motivo}")
        return 1
    sid = grava(est, pre, args.data, f"{pathlib.Path(args.estoque).name} + "
                                     f"{pathlib.Path(args.preco).name}")
    print(f"\n  snapshot {sid} gravado em {args.data}: "
          f"{len(est):,} linha(s) de estoque, {len(pre):,} de preco")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
