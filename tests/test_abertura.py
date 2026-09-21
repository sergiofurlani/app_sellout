"""O saldo de abertura congela o passado — e só ele."""

import csv
from datetime import date

import openpyxl
import pytest

from sellout.db import abertura, movimentos


def planilha(tmp_path, blocos):
    """blocos: [(titulo, [(codigo, descricao, estoque_inicial)])]"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Masculino"
    r = 3                      # o cabeçalho de cada bloco é a linha 3 do primeiro
    for titulo, produtos in blocos:
        ws.cell(r, 2, "Código")
        ws.cell(r, 3, titulo)
        ws.cell(r, 10, "Estoque inicial")
        r += 1
        for codigo, desc, ei in produtos:
            ws.cell(r, 2, codigo)
            ws.cell(r, 3, desc)
            ws.cell(r, 10, ei)
            r += 1
        ws.cell(r, 3, "Subtotal")
        r += 2
    wb.create_sheet("Feminino")
    caminho = tmp_path / "p.xlsx"
    wb.save(caminho)
    return str(caminho)


def test_colecao_reconstruida_nao_e_congelada(tmp_path):
    arq = planilha(tmp_path, [
        ("MASCULINO - AW26", [("330001", "CALCA", 100)]),
        ("MASCULINO - SS25", [("320001", "CAMISA", 40)]),
    ])
    linhas, resumo, _ = abertura.ler(arq)
    codigos = {l["codigo"] for l in linhas}
    assert codigos == {"320001"}, "só a coleção antiga é congelada"
    assert resumo["pecas congeladas"] == 40
    assert resumo["pecas reconstruidas"] == 100


def test_codigo_nos_dois_lados_fica_inteiro_com_o_erp(tmp_path):
    """Congelar a metade antiga criaria estoque do nada quando o ERP somasse
    a parte dele — o mesmo código seria contado duas vezes."""
    arq = planilha(tmp_path, [
        ("MASCULINO - AW26", [("330043", "CALCA MB", 60)]),
        ("MASCULINO - SS25", [("330043", "CALCA MB", 25)]),
    ])
    linhas, resumo, avisos = abertura.ler(arq)
    assert linhas == []
    assert resumo["pecas reconstruidas"] == 85
    assert any("330043" in a for a in avisos)


def test_duas_linhas_da_mesma_colecao_somam(tmp_path):
    arq = planilha(tmp_path, [
        ("MASCULINO - SS25", [("320001", "CAMISA AZUL", 30),
                              ("320001", "CAMISA DEMAIS CORES", 12)]),
    ])
    linhas, _resumo, avisos = abertura.ler(arq)
    assert len(linhas) == 1
    assert linhas[0]["qtd"] == 42
    assert any("2 linhas" in a for a in avisos)


def test_zerado_nao_vira_linha(tmp_path):
    arq = planilha(tmp_path, [("MASCULINO - SS25", [("320001", "CAMISA", 0)])])
    linhas, resumo, _ = abertura.ler(arq)
    assert linhas == []
    assert resumo["zerados, nao congelados"] == 1


def csv_movimentos(tmp_path, linhas):
    caminho = tmp_path / "mov.csv"
    campos = ["data", "tipo", "origem", "destino", "codigo", "cod_cor", "cor",
              "tamanho", "quant", "valor", "nota", "romaneio"]
    with open(caminho, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=campos, delimiter=";")
        w.writeheader()
        for l in linhas:
            w.writerow({c: l.get(c, "") for c in campos})
    return str(caminho)


def test_entrada_e_saida_mudam_de_sinal_e_de_filial(tmp_path):
    arq = csv_movimentos(tmp_path, [
        {"data": "2026-01-06", "tipo": "entrada", "origem": "ELENA ES",
         "destino": "IGUATEMI", "codigo": "330001", "cod_cor": "08",
         "tamanho": "38", "quant": "5", "romaneio": "111"},
        {"data": "2026-01-07", "tipo": "saida", "origem": "EGREY JDS",
         "destino": "ELENA ES", "codigo": "330001", "cod_cor": "08",
         "tamanho": "38", "quant": "2", "romaneio": "112"},
    ])
    linhas, resumo, _ = movimentos.ler(arq)
    entrada = [l for l in linhas if l["tipo"] == "transferencia"][0]
    saida = [l for l in linhas if l["tipo"] == "devolucao"][0]
    assert entrada["qtd"] == 5 and entrada["filial"] == "IGUATEMI"
    assert saida["qtd"] == -2 and saida["filial"] == "EGREY JDS"
    assert entrada["data"] == date(2026, 1, 6)
    assert resumo["pecas transferencia"] == 5


def test_realocacao_entra_dos_dois_lados_e_soma_zero(tmp_path):
    arq = csv_movimentos(tmp_path, [
        {"data": "2026-03-02", "tipo": "realocacao", "origem": "EGREY JDS",
         "destino": "IGUATEMI", "codigo": "330001", "cod_cor": "08",
         "quant": "3", "romaneio": "222"},
    ])
    linhas, _resumo, _ = movimentos.ler(arq)
    assert len(linhas) == 2
    assert sum(l["qtd"] for l in linhas) == 0
    assert {l["filial"] for l in linhas} == {"EGREY JDS", "IGUATEMI"}


def test_fora_do_sellout_nao_entra(tmp_path):
    arq = csv_movimentos(tmp_path, [
        {"data": "2026-03-02", "tipo": "fora", "origem": "ELENA ES",
         "destino": "OUTLET", "codigo": "330001", "quant": "9"},
    ])
    linhas, resumo, _ = movimentos.ler(arq)
    assert linhas == []
    assert resumo["fora do sellout"] == 1


def test_linha_sem_data_avisa_em_vez_de_quebrar(tmp_path):
    arq = csv_movimentos(tmp_path, [
        {"data": "", "tipo": "entrada", "origem": "ELENA ES", "destino": "IGUATEMI",
         "codigo": "330001", "quant": "4"},
    ])
    linhas, _resumo, avisos = movimentos.ler(arq)
    assert linhas == []
    assert avisos and "sem data" in avisos[0]


# ---------------------------------------------------------------- fontes ERP

def xml_estoque(tmp_path, linhas):
    cab = ["Filial", "Código", "Descrição", "CODIGO", "Cor", "tamanho",
           "ESTOQUE ATUAL"]
    def celula(v):
        return f'<Cell><Data ss:Type="String">{v}</Data></Cell>'
    corpo = "".join("<Row>" + "".join(celula(v) for v in l) + "</Row>"
                    for l in [cab] + linhas)
    caminho = tmp_path / "e.xls"
    caminho.write_text(
        '<?xml version="1.0"?><Workbook '
        'xmlns="urn:schemas-microsoft-com:office:spreadsheet" '
        'xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet">'
        f'<Worksheet ss:Name="Plan1"><Table>{corpo}</Table></Worksheet></Workbook>',
        encoding="utf-8")
    return str(caminho)


def test_estoque_traduz_a_filial_e_separa_a_cor_do_produto(tmp_path):
    from coletor import fontes_erp
    arq = xml_estoque(tmp_path, [
        ["EGREY IGUATEMI", "100001", "T-SHIRT", "0002", "PRETO", "G", "7"],
        ["EGREY JARDINS", "100001", "T-SHIRT", "0002", "PRETO", "M", "3"],
    ])
    linhas, resumo, avisos = fontes_erp.le_estoque(arq)
    assert avisos == []
    assert {l["filial"] for l in linhas} == {"IGUATEMI", "EGREY JDS"}
    assert linhas[0]["codigo"] == "100001" and linhas[0]["codigo_cor"] == "0002"
    assert resumo["pecas"] == 10


def test_filial_fora_do_mapa_avisa_mas_nao_some(tmp_path):
    """Perder linha em silêncio por causa de um nome novo de filial seria o
    mesmo defeito da D4, que já custou 3.000 peças de estoque sumido."""
    from coletor import fontes_erp
    arq = xml_estoque(tmp_path, [["EGREY NOVA LOJA", "100001", "X", "0002", "PRETO", "G", "5"]])
    linhas, resumo, avisos = fontes_erp.le_estoque(arq)
    assert len(linhas) == 1 and linhas[0]["filial"] == "EGREY NOVA LOJA"
    assert avisos and "EGREY NOVA LOJA" in avisos[0]
    assert resumo["pecas"] == 5


def test_confere_acha_o_estoque_sem_preco(tmp_path):
    from coletor import fontes_erp
    est = [{"codigo": "1", "codigo_cor": "0002", "tamanho": "G", "qtd": 9},
           {"codigo": "2", "codigo_cor": "0002", "tamanho": "G", "qtd": 4}]
    pre = [{"codigo": "1", "codigo_cor": "0002", "tamanho": "G", "preco": 100.0}]
    c = fontes_erp.confere(est, pre)
    assert c["sem_preco"] == 1 and c["pecas_sem_preco"] == 4


def test_le_os_dois_formatos_de_export(tmp_path):
    """O ERP manda `.xls` que é XML e `.xlsx` de verdade; e a coluna da cor
    mudou de `CODIGO` para `CODIGO_COR` no meio do caminho."""
    from coletor import fontes_erp
    import openpyxl as _o
    velho = xml_estoque(tmp_path, [
        ["EGREY IGUATEMI", "100001", "T-SHIRT", "0002", "PRETO", "G", "7"]])

    wb = _o.Workbook(); ws = wb.active
    ws.append(["Filial", "Código", "Descrição", "CODIGO_COR", "Cor", "tamanho",
               "ESTOQUE ATUAL"])
    ws.append(["EGREY IGUATEMI", "100001", "T-SHIRT", "0002", "PRETO", "G", 7])
    novo = tmp_path / "novo.xlsx"; wb.save(novo)

    a, _ra, _ = fontes_erp.le_estoque(velho)
    b, _rb, _ = fontes_erp.le_estoque(str(novo))
    assert a == b, "os dois formatos têm que dar a mesma linha"
    assert b[0]["codigo_cor"] == "0002"


def test_codigo_fora_do_sellout_nao_entra_e_e_contado(tmp_path):
    """O SW4-WH-FSC carregava 1.000 das 10.083 peças sem ser peça de venda.
    Sair calado seria trocar um erro por outro — ele sai contado."""
    from coletor import fontes_erp
    arq = xml_estoque(tmp_path, [
        ["EGREY IGUATEMI", "SW4-WH-FSC", "EMBALAGEM", "000", "UNICA", "U", "1000"],
        ["EGREY IGUATEMI", "100001", "T-SHIRT", "0002", "PRETO", "G", "7"],
    ])
    linhas, resumo, _ = fontes_erp.le_estoque(arq)
    assert [l["codigo"] for l in linhas] == ["100001"]
    assert resumo["pecas"] == 7
    assert resumo["fora do sellout"] == 1 and resumo["pecas fora"] == 1000
