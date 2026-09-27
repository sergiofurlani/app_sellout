"""A exportação do ERP contra o banco.

É a conferência que decide se a extração é fiel — e é a única das três em que a
tolerância é zero. Por isso as funções que comparam são testadas sem banco.
"""

import openpyxl
import pytest

from sellout.db import confere_erp


def chave(cod="332004", cor="0002", fil="IGUATEMI"):
    return (cod, cor, fil)


def test_iguais_na_peca_nao_entram_em_divergencia():
    r = confere_erp.compara({chave(): 3.0}, {chave(): 3.0})
    assert r["iguais"] == [chave()] and r["difere"] == []
    assert confere_erp.fiel(r)


def test_uma_peca_de_diferenca_e_achado():
    """**O oposto da conferência contra a planilha.** Lá uma peça sai de data de
    corte; aqui é a mesma semana, a mesma fonte e o mesmo grão."""
    r = confere_erp.compara({chave(): 3.0}, {chave(): 2.0})
    assert r["difere"] == [{"chave": chave(), "erp": 3.0, "banco": 2.0, "dif": 1.0}]
    assert not confere_erp.fiel(r)


def test_chave_de_um_lado_so_nao_vira_diferenca_de_numero():
    """Linha que um dos dois não viu é outra coisa que número errado. Somar zero
    do outro lado misturaria as duas — erro já corrigido na outra conferência."""
    r = confere_erp.compara({chave(): 2.0}, {chave(cor="0008"): 2.0})
    assert [x["chave"] for x in r["so_erp"]] == [chave()]
    assert [x["chave"] for x in r["so_banco"]] == [chave(cor="0008")]
    assert r["difere"] == []
    assert not confere_erp.fiel(r)


def test_devolucao_negativa_e_comparada_como_qualquer_numero():
    """A exportação traz devolução com sinal negativo, e a `venda` também."""
    r = confere_erp.compara({chave(): -1.0}, {chave(): -1.0})
    assert confere_erp.fiel(r)


def test_o_sinal_trocado_nao_passa():
    r = confere_erp.compara({chave(): -1.0}, {chave(): 1.0})
    assert r["difere"][0]["dif"] == -2.0


def test_a_diferenca_e_agrupada_por_filial():
    """Uma filial inteira faltando é o defeito que a soma total esconde: 3 a
    mais numa e 3 a menos noutra dão zero no total."""
    r = confere_erp.compara(
        {chave(fil="SITE"): 5.0, chave(fil="IGUATEMI"): 2.0},
        {chave(fil="SITE"): 2.0, chave(fil="IGUATEMI"): 5.0})
    pf = confere_erp.por_filial(r)
    assert pf["SITE"]["dif"] == 3.0
    assert pf["IGUATEMI"]["dif"] == -3.0
    assert r["total_erp"] == r["total_banco"] == 7.0, "no total desaparece"


def test_filial_que_so_existe_num_lado_aparece_na_tabela():
    r = confere_erp.compara({chave(fil="OUTLET"): 4.0}, {})
    pf = confere_erp.por_filial(r)
    assert pf["OUTLET"] == {"erp": 4.0, "banco": 0.0, "dif": 4.0, "chaves": 1}


def test_os_totais_somam_os_dois_lados():
    r = confere_erp.compara({chave(): 3.0, chave(cor="0008"): 1.0}, {chave(): 2.0})
    assert r["total_erp"] == 4.0 and r["total_banco"] == 2.0


# ------------------------------------------------- leitura da exportacao

CABECALHO = ("Filial", "Código", "Descrição", "Codigo_cor", "Cor", "Tam",
             "Qtde", "Valor")


def export(tmp_path, *linhas, cabecalho=CABECALHO):
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, t in enumerate(cabecalho, 1):
        ws.cell(1, i, t)
    for r, linha in enumerate(linhas, 2):
        for i, v in enumerate(linha, 1):
            ws.cell(r, i, v)
    caminho = tmp_path / "v.xlsx"
    wb.save(caminho)
    return str(caminho)


def linha(fil="SITE", cod="332004", cor="0002", tam="38", qtd=1, valor=100):
    return (fil, cod, "BLUSA", cor, "PRETO", tam, qtd, valor)


def test_os_tamanhos_somam_no_produto_e_cor(tmp_path):
    """O sellout é por produto e cor; troca de tamanho dentro do produto não
    muda nada do que está sendo medido."""
    d, _ = confere_erp.do_export(export(tmp_path, linha(tam="38"), linha(tam="40")))
    assert d == {("332004", "0002", "SITE"): 2.0}


def test_a_devolucao_negativa_e_somada_junto(tmp_path):
    d, _ = confere_erp.do_export(export(tmp_path, linha(qtd=3), linha(qtd=-1)))
    assert d[("332004", "0002", "SITE")] == 2.0


def test_as_colunas_sao_achadas_pelo_cabecalho(tmp_path):
    """Relatório de ERP ganha coluna sem avisar — foi o que a "Consignado Real"
    fez no meio da planilha."""
    cab = ("Valor", "Qtde", "Tam", "Cor", "Codigo_cor", "Descrição", "Código",
           "Filial")
    c = export(tmp_path, (100, 2, "38", "PRETO", "0002", "BLUSA", "332004", "SITE"),
               cabecalho=cab)
    d, _ = confere_erp.do_export(c)
    assert d == {("332004", "0002", "SITE"): 2.0}


def test_exportacao_sem_as_colunas_necessarias_para_com_erro(tmp_path):
    """Ler uma exportação de outro relatório daria totais plausíveis e errados."""
    c = export(tmp_path, ("SITE", "BLUSA"), cabecalho=("Filial", "Descrição"))
    with pytest.raises(ValueError, match="nao tem as colunas"):
        confere_erp.do_export(c)


def test_quantidade_nao_numerica_e_avisada_e_nao_somada(tmp_path):
    """Linha de total ou de rodapé no meio da exportação não pode entrar na
    conta calada."""
    d, avisos = confere_erp.do_export(
        export(tmp_path, linha(qtd=2), linha(qtd="total")))
    assert d[("332004", "0002", "SITE")] == 2.0
    assert len(avisos) == 1 and "nao numerica" in avisos[0]


def test_linha_sem_codigo_e_ignorada(tmp_path):
    d, _ = confere_erp.do_export(export(tmp_path, linha(cod=""), linha()))
    assert len(d) == 1


# ------------------------------- a mesma loja com dois nomes (27/09)

def test_jardins_do_relatorio_e_egrey_jds_do_banco():
    """**O achado da primeira corrida.** 211 das 219 divergências eram uma loja
    com dois nomes: o relatório escreve JARDINS, a API devolve o código da
    filial. Sem isto a conferência acusa a loja inteira dos dois lados."""
    assert confere_erp.filial_do_export("JARDINS") == "EGREY JDS"
    assert confere_erp.filial_do_export(" jardins ") == "EGREY JDS"


def test_filial_sem_apelido_passa_como_esta():
    assert confere_erp.filial_do_export("IGUATEMI") == "IGUATEMI"
    assert confere_erp.filial_do_export("SITE") == "SITE"


def test_o_apelido_e_aplicado_na_leitura(tmp_path):
    d, _ = confere_erp.do_export(export(tmp_path, linha(fil="JARDINS")))
    assert list(d) == [("332004", "0002", "EGREY JDS")]


def test_nome_de_um_lado_so_e_denunciado():
    """O sintoma do apelido faltando tem de gritar — vale para a loja nova que
    ninguém mapeou ainda, que produziria centenas de divergências falsas."""
    s = confere_erp.so_de_um_lado({chave(fil="OUTLET"): 1.0},
                                  {chave(fil="EGREY JDS"): 1.0})
    assert s == {"export": ["OUTLET"], "banco": ["EGREY JDS"]}


def test_filiais_que_casam_nao_sao_denunciadas():
    s = confere_erp.so_de_um_lado({chave(fil="SITE"): 1.0}, {chave(fil="SITE"): 2.0})
    assert s == {"export": [], "banco": []}
