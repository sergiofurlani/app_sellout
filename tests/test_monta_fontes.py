"""A ponte que escreve as abas de origem na planilha."""

import csv

import openpyxl
import pytest

from coletor import monta_fontes


def planilha(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Masculino"
    ws.cell(3, 2, "Código"); ws.cell(3, 3, "MASCULINO - SS27")
    ws.cell(3, 10, "Estoque inicial")
    ws.cell(4, 2, "100001"); ws.cell(4, 3, "T-SHIRT"); ws.cell(4, 10, 12)
    for nome in ("Feminino", "Estoque", "Vendas", "Producao", "Preco", "Produtos"):
        wb.create_sheet(nome)
    wb["Producao"].append(["CODIGO", "COR", "TAMANHO", "QUANTIDADE"])
    wb["Producao"].append(["100001", "PRETO", "G", 5])
    caminho = tmp_path / "g.xlsx"
    wb.save(caminho)
    return str(caminho)


def test_escreve_as_tres_abas_e_nao_toca_nas_outras(tmp_path):
    origem = planilha(tmp_path)
    saida = str(tmp_path / "saida.xlsx")
    monta_fontes.monta(
        origem,
        [{"filial": "IGUATEMI", "codigo": "100001", "codigo_cor": "0002",
          "cor": "PRETO", "tamanho": "G", "qtd": 7}],
        [{"codigo": "100001", "codigo_cor": "0002", "tamanho": "G", "preco": 699.0}],
        [{"filial": "IGUATEMI", "codigo": "100001", "codigo_cor": "0002",
          "cor": "PRETO", "tamanho": "G", "qtd": 2, "valor": 1398.0}],
        saida)

    wb = openpyxl.load_workbook(saida)
    assert [c.value for c in wb["Estoque"][1]] == monta_fontes.CABECALHOS["Estoque"]
    assert wb["Estoque"]["A2"].value == "IGUATEMI"
    assert wb["Vendas"]["G2"].value == 2
    assert wb["Preco"]["F2"].value == 699.0
    # produção intacta: ela alimentava o Estoque inicial, que hoje vem do 106
    assert wb["Producao"]["A2"].value == "100001"
    # a aba de trabalho não foi tocada
    assert wb["Masculino"]["C4"].value == "T-SHIRT"
    assert wb["Masculino"]["J4"].value == 12


def test_aba_que_falta_para_em_vez_de_escrever_pela_metade(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Masculino"
    caminho = tmp_path / "incompleta.xlsx"
    wb.save(caminho)
    with pytest.raises(ValueError, match="Estoque"):
        monta_fontes.monta(str(caminho), [], [], [], str(tmp_path / "x.xlsx"))


def test_le_vendas_do_csv_do_coletor(tmp_path):
    caminho = tmp_path / "v.csv"
    with open(caminho, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["FILIAL", "CODIGO", "DESCRICAO", "CODIGO COR", "COR",
                    "TAMANHO", "QTDE", "VALOR"])
        w.writerow(["IGUATEMI", "334128", "", "0002", "PRETO", "38", "2", "1398.00"])
        w.writerow(["IGUATEMI", "", "", "", "", "", "9", "0"])   # sem código: fora
    linhas = monta_fontes.le_vendas(str(caminho))
    assert len(linhas) == 1
    assert linhas[0]["qtd"] == 2 and linhas[0]["valor"] == 1398.0


def _com_valor_guardado(caminho, formula="I4/J4", valor="0.5"):
    """Injeta o resultado guardado de uma fórmula, como o Excel faz.

    O openpyxl não sabe gravar isso — ele nunca calcula —, então o cenário que
    quebrou em 21/09 não podia ser montado com as ferramentas do projeto. Sem
    este remendo de três linhas o defeito não tem teste.
    """
    import re
    import zipfile

    with zipfile.ZipFile(caminho) as zf:
        itens = [(i, zf.read(i.filename)) for i in zf.infolist()]
    with zipfile.ZipFile(caminho, "w", zipfile.ZIP_DEFLATED) as saida:
        for item, dados in itens:
            if item.filename.startswith("xl/worksheets/"):
                dados = re.sub(rb"<f>" + formula.encode() + rb"</f>",
                               b"<f>" + formula.encode() + b"</f><v>"
                               + valor.encode() + b"</v>", dados)
            saida.writestr(item, dados)
    return caminho


def test_o_valor_guardado_da_formula_sobrevive(tmp_path):
    """A coluna `Sellout 14/09` saiu em branco na planilha de 21/09.

    O openpyxl não calcula fórmula e descarta o resultado que o Excel guardou.
    O app congela a coluna da semana passada lendo justamente esse valor
    (`data_only=True`): achou None e escreveu vazio. Uma semana de histórico
    perdida por causa de um `wb.save()`.
    """
    origem = planilha(tmp_path)
    wb = openpyxl.load_workbook(origem)
    wb["Masculino"].cell(4, 11, "=I4/J4")
    wb.save(origem)
    _com_valor_guardado(origem)

    assert openpyxl.load_workbook(origem, data_only=True)["Masculino"].cell(4, 11).value == 0.5

    saida = str(tmp_path / "saida.xlsx")
    monta_fontes.monta(origem, [], [], [], saida)

    lido = openpyxl.load_workbook(saida, data_only=True)["Masculino"].cell(4, 11).value
    assert lido == 0.5, "o app congela a coluna da semana passada com este valor"
    formula = openpyxl.load_workbook(saida)["Masculino"].cell(4, 11).value
    assert formula == "=I4/J4", "a fórmula tem que continuar lá para a próxima rodada"


def test_a_contagem_acusa_a_perda(tmp_path):
    """A guarda que faltava: o vermelho era conferido antes e depois, o valor
    guardado não era — e era ele que sumia, calado."""
    origem = planilha(tmp_path)
    wb = openpyxl.load_workbook(origem)
    wb["Masculino"].cell(4, 11, "=I4/J4")
    wb.save(origem)
    _com_valor_guardado(origem)
    assert monta_fontes.conta_formulas_com_valor(origem) == 1

    # o caminho antigo, que salvava a planilha inteira pelo openpyxl
    regravado = str(tmp_path / "regravado.xlsx")
    openpyxl.load_workbook(origem).save(regravado)
    assert monta_fontes.conta_formulas_com_valor(regravado) == 0
