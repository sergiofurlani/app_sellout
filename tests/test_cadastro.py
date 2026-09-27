"""Semeadura das dimensões: coleção do ERP, linha comercial da planilha."""

import pytest

from sellout.db import cadastro, conexao

precisa_banco = pytest.mark.skipif(
    not conexao.disponivel(), reason="sem DATABASE_URL/psycopg")

CABECALHO = ("codigo;interno;colecao;subcolecao;sigla;referencia;descricao;"
             "tipo;grupo;departamento;marca;divisao;categoria;status;grade;"
             "fornecedor;cadastro\n")


def csv_de(tmp_path, *linhas):
    caminho = tmp_path / "produtos-erp.csv"
    caminho.write_text(CABECALHO + "".join(linhas), encoding="utf-8-sig")
    return str(caminho)


def linha(codigo="334128", colecao="VERÃO", sub="2027", sigla="SS27",
          descricao="BLUSA", divisao="FEMININO"):
    return (f"{codigo};1;{colecao};{sub};{sigla};ref;{descricao};"
            f"tipo;grupo;dep;marca;{divisao};cat;ativo;grade;forn;2026-01-01\n")


def test_a_colecao_gravada_e_a_sigla():
    """SS27, não VERÃO/2027. É assim que o negócio fala e é o que os blocos da
    planilha usam; guardar a forma crua obrigaria cada leitor a remontar a
    sigla, e a regra ficaria repetida em cada consulta."""
    import tempfile, pathlib
    with tempfile.TemporaryDirectory() as d:
        c = csv_de(pathlib.Path(d), linha())
        produtos, _ = cadastro.do_cadastro(c)
    assert produtos[0]["colecao"] == "SS27"


def test_perene_vira_colecao_dos_classicos(tmp_path):
    """PERENE não tem estação e por isso não gera sigla — mas é o universo dos
    Clássicos e precisa de nome no banco (D11)."""
    c = csv_de(tmp_path, linha(colecao="PERENE", sub="", sigla=""))
    produtos, _ = cadastro.do_cadastro(c)
    assert produtos[0]["colecao"] == "PERENE"


def test_colecao_fora_do_mapa_fica_sem_sigla_e_e_reportada(tmp_path):
    """Produto sem agrupamento some de qualquer recorte por coleção. Deixar
    isso calado é o mesmo defeito que travou a comparação em 27/09."""
    c = csv_de(tmp_path, linha(colecao="INDEFINIDO", sub="", sigla=""))
    produtos, motivos = cadastro.do_cadastro(c)
    assert produtos[0]["colecao"] == ""
    assert motivos == {"INDEFINIDO": 1}


def test_o_ano_vem_da_subcolecao(tmp_path):
    c = csv_de(tmp_path, linha(codigo="219054", colecao="INVERNO",
                               sub="2026", sigla="AW26"))
    produtos, _ = cadastro.do_cadastro(c)
    assert produtos[0]["colecao"] == "AW26"


def test_le_os_campos_que_a_tabela_produto_tem(tmp_path):
    c = csv_de(tmp_path, linha())
    (p,), _ = cadastro.do_cadastro(c)
    assert set(p) == {"codigo", "colecao", "descricao", "divisao",
                      "departamento", "grupo", "marca", "grade"}


def test_codigo_vazio_nao_entra(tmp_path):
    c = csv_de(tmp_path, linha(codigo=""), linha())
    produtos, _ = cadastro.do_cadastro(c)
    assert len(produtos) == 1


@precisa_banco
def test_gravar_cadastro_e_idempotente(tmp_path):
    c = csv_de(tmp_path, linha(codigo="999999"))
    produtos, _ = cadastro.do_cadastro(c)
    assert cadastro.grava_cadastro(produtos) == 1
    assert cadastro.grava_cadastro(produtos) == 1


# ------------------------------------------- a classificacao atual da planilha

def onde(codigo="334128", bloco="SS27", colecao="SS27", linha=""):
    return {codigo: {"bloco": bloco, "aba": "Feminino",
                     "colecao": colecao, "linha": linha}}


def test_bloco_de_colecao_e_comparado_com_o_erp():
    d = cadastro.divergencias(onde(), {"334128": "SS27"})
    assert d["iguais"] == ["334128"] and d["difere"] == []


def test_a_divergencia_diz_os_dois_lados():
    """Saber que difere não basta — é a lista de trabalho para arrumar o
    cadastro, e ela precisa dizer o que está de cada lado."""
    d = cadastro.divergencias(onde(), {"334128": "AW26"})
    assert d["difere"] == [{"codigo": "334128", "planilha": "SS27", "erp": "AW26"}]


def test_produto_sem_colecao_no_cadastro_e_separado_de_quem_nao_esta_nele():
    """São dois problemas diferentes: um produto existe e está sem
    classificação; o outro nem está cadastrado. Misturar os dois manda a
    pessoa procurar no lugar errado."""
    d = cadastro.divergencias(
        {**onde("111111"), **onde("222222")},
        {"111111": ""})
    assert [x["codigo"] for x in d["sem_colecao_erp"]] == ["111111"]
    assert [x["codigo"] for x in d["so_planilha"]] == ["222222"]


def test_bloco_de_linha_comercial_nao_entra_na_comparacao():
    """HOME e PIMA não existem no ERP (D11). Cobrá-los do cadastro criaria uma
    divergência que nunca some."""
    d = cadastro.divergencias(onde(bloco="HOME", colecao="", linha="HOME"), {})
    assert d == {"iguais": [], "difere": [], "so_planilha": [],
                 "sem_colecao_erp": []}


def planilha_com_blocos(tmp_path):
    """Duas abas, três blocos: dois de coleção e um de linha comercial."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Feminino"
    def bloco(r, titulo, codigos):
        ws.cell(r, 2, "Código"); ws.cell(r, 3, titulo)
        for i, c in enumerate(codigos, 1):
            ws.cell(r + i, 2, c)
        return r + len(codigos) + 2
    r = bloco(3, "FEMININO - SS27", ["334128", "334129"])
    r = bloco(r, "FEMININO - HOME", ["140000"])
    ws2 = wb.create_sheet("Masculino")
    ws2.cell(3, 2, "Código"); ws2.cell(3, 3, "MASCULINO - AW26")
    ws2.cell(4, 2, "219054")
    caminho = tmp_path / "g.xlsx"
    wb.save(caminho)
    return str(caminho)


def test_a_planilha_registra_todo_bloco_nao_so_a_linha_comercial(tmp_path):
    """**O ajuste de 27/09.** A primeira versão guardava só HOME/PIMA e
    descartava AW26 e SS27, supondo que coleção era assunto exclusivo do ERP.
    Mas é a classificação atual da planilha que permite auditar o cadastro
    enquanto ele é arrumado — sem ela não há contra o quê comparar."""
    onde = cadastro.da_planilha(planilha_com_blocos(tmp_path))
    assert set(onde) == {"334128", "334129", "140000", "219054"}
    assert onde["334128"]["colecao"] == "SS27"
    assert onde["219054"]["colecao"] == "AW26"


def test_o_bloco_de_colecao_nao_vira_linha_comercial(tmp_path):
    """SS27 é coleção e vem do ERP. Gravá-lo como linha comercial misturaria
    dimensão nossa com dimensão do ERP na mesma coluna."""
    onde = cadastro.da_planilha(planilha_com_blocos(tmp_path))
    assert onde["334128"]["linha"] == ""
    assert onde["140000"]["linha"] == "HOME"
    assert onde["140000"]["colecao"] == ""


def test_a_aba_de_origem_fica_registrada(tmp_path):
    onde = cadastro.da_planilha(planilha_com_blocos(tmp_path))
    assert onde["219054"]["aba"] == "Masculino"
    assert onde["334128"]["aba"] == "Feminino"


# ----------------------------------------------------- excecoes de colecao

def produto(codigo="212006", colecao="SS24"):
    return {"codigo": codigo, "colecao": colecao, "descricao": "", "divisao": "",
            "departamento": "", "grupo": "", "marca": "", "grade": ""}


def test_a_excecao_sobrepoe_a_colecao_do_erp():
    """Decisão do negócio em 27/09: nestes 21 códigos vale a planilha."""
    ps = [produto()]
    aplicadas, _ = cadastro.aplica_excecoes(ps, {"212006": "SS25"})
    assert ps[0]["colecao"] == "SS25"
    assert aplicadas == [{"codigo": "212006", "erp": "SS24", "decidida": "SS25"}]


def test_produto_fora_da_lista_nao_e_tocado():
    ps = [produto(codigo="999999", colecao="AW26")]
    cadastro.aplica_excecoes(ps, {"212006": "SS25"})
    assert ps[0]["colecao"] == "AW26"


def test_excecao_que_ja_concorda_com_o_cadastro_e_apontada():
    """**A lista precisa encolher.** Quando o cadastro for corrigido, a exceção
    vira linha morta — e exceção que ninguém tira acaba escondendo uma decisão
    que já não vale."""
    ps = [produto(colecao="SS25")]
    aplicadas, redundantes = cadastro.aplica_excecoes(ps, {"212006": "SS25"})
    assert aplicadas == []
    assert redundantes == ["212006"]


def test_o_arquivo_de_excecoes_e_lido_do_repositorio():
    """Versionado, com motivo e data: ajuste de dimensão sem rastro é o tipo de
    coisa que ninguém explica seis meses depois."""
    e = cadastro.le_excecoes()
    assert e["212006"] == "SS25"
    assert e["338005"] == "AW26", "o ERP dizia SS27 neste"
    assert len(e) == 21


def test_sem_arquivo_nao_ha_excecao(tmp_path):
    assert cadastro.le_excecoes(tmp_path / "nao-existe.csv") == {}
